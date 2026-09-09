"""Разбор XLSX в `Document`. План М7.

Сегмент — одна непустая ячейка листа. Якорь — ``("cell", имя_листа,
номер_строки, номер_столбца)``, оба индекса 1-based, как у ``openpyxl``
(``Cell.row``/``Cell.column``): рендер обращается к ячейке напрямую через
``worksheet.cell(row=, column=)``, без пересчёта индекса и без привязки к
смещению в линейном тексте — иначе не собрать ни подсветку, ни отчёт (см.
AGENTS.md, «Якоря сегментов»).

Ячейка-формула не сканируется по своему коду: ``=B2*1.2`` — это программа,
а не данные, и поиск в ней ИНН или ФИО даёт только ложные срабатывания на
адресах ячеек и именах функций. Вместо кода сегмент строится из
закэшированного результата вычисления — того, что реально видит человек,
открывший файл в Excel/LibreOffice (см. ``_display_text``, `iter_cells`).
Формула без кэша (файл создан программно и ни разу не открывался в офисном
приложении) не даёт сегмента вовсе: сканировать нечего, а нормализованный
номер ячейки текстом не является. Это не дыра в защите: формула, зависящая
от замаскированной ячейки, обезвреживается отдельно и безусловно в
``render/xlsx_redact.py`` — вне зависимости от того, что нашёл детектор по
тексту.
"""

from __future__ import annotations

import datetime
import pathlib
from collections.abc import Iterator

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook

from masker.model import Anchor, Document, Segment

#: Дискриминатор якоря — единственная часть XLSX, которую разбирает T1.1-аналог
#: для этого формата: обычные ячейки листов. Сводные таблицы, комментарии и
#: колонтитулы в этот разбор не входят.
PART_CELL = "cell"

XlsxLocator = tuple[str | int, ...]


def _display_text(value: object) -> str:
    """Текст ячейки так, как его видит человек, а не как хранит Python.

    Формат — не бухгалтерский (``number_format`` ячейки не учитывается,
    это потребовало бы повторной реализации форматирования Excel), а
    достаточный для детектора: пробел не отличим от отсутствия значения,
    дата приведена к обычной русской числовой записи ``ДД.ММ.ГГГГ``.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "ИСТИНА" if value else "ЛОЖЬ"
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def iter_cells(
    formula_workbook: Workbook, value_workbook: Workbook
) -> Iterator[tuple[XlsxLocator, str]]:
    """Пройти листы обеих раскрытых копий книги и отдать текст непустых ячеек.

    ``formula_workbook`` открыта без ``data_only`` — по ней определяется,
    формула ли ячейка (``cell.data_type == "f"``). ``value_workbook``
    открыта с ``data_only=True`` — из неё берётся текст: для обычных ячеек
    оба значения совпадают, а для формул вторая книга отдаёт кэш
    вычисления вместо кода. Два экземпляра книги, а не один с постфактум
    переключением режима, — ``openpyxl`` фиксирует ``data_only`` на этапе
    разбора файла и не даёт посмотреть на одну и ту же ячейку и так, и так.
    """
    for sheet_name in formula_workbook.sheetnames:
        formula_sheet = formula_workbook[sheet_name]
        value_sheet = value_workbook[sheet_name]
        for formula_row, value_row in zip(
            formula_sheet.iter_rows(), value_sheet.iter_rows(), strict=True
        ):
            for formula_cell, value_cell in zip(formula_row, value_row, strict=True):
                is_formula = formula_cell.data_type == "f"
                text = _display_text(value_cell.value if is_formula else formula_cell.value)
                if not text.strip():
                    continue
                locator: XlsxLocator = (
                    PART_CELL,
                    sheet_name,
                    formula_cell.row,
                    formula_cell.column,
                )
                yield locator, text


def _anchor_label(locator: XlsxLocator) -> str:
    _, sheet_name, row, col = locator
    assert isinstance(sheet_name, str)
    assert isinstance(row, int)
    assert isinstance(col, int)
    return f"лист «{sheet_name}», {get_column_letter(col)}{row}"


def ingest_xlsx(path: str | pathlib.Path) -> Document:
    """Разобрать `.xlsx` в `Document` с сегментами и якорями `(лист, строка, столбец)`."""
    path = pathlib.Path(path)
    formula_workbook = load_workbook(str(path), data_only=False)
    value_workbook = load_workbook(str(path), data_only=True)

    segments: list[Segment] = []
    for locator, text in iter_cells(formula_workbook, value_workbook):
        segments.append(
            Segment(
                text=text,
                anchor=Anchor(fmt="xlsx", locator=locator, label=_anchor_label(locator)),
                order=len(segments),
            )
        )

    props = formula_workbook.properties
    meta = {
        key: value
        for key, value in (
            ("author", props.creator),
            ("last_modified_by", props.lastModifiedBy),
            ("title", props.title),
            ("subject", props.subject),
            ("comments", props.description),
            ("category", props.category),
            ("keywords", props.keywords),
        )
        if value
    }

    formula_workbook.close()
    value_workbook.close()
    return Document(path=str(path), fmt="xlsx", segments=segments, meta=meta)


def resolve_anchor(workbook: Workbook, locator: XlsxLocator) -> Cell | None:
    """Найти ячейку XLSX по локатору сегмента в открытой копии книги."""
    if len(locator) != 4 or locator[0] != PART_CELL:
        return None
    _, sheet_name, row, col = locator
    if not isinstance(sheet_name, str) or not isinstance(row, int) or not isinstance(col, int):
        return None
    if sheet_name not in workbook.sheetnames:
        return None
    if row < 1 or col < 1:
        return None
    return workbook[sheet_name].cell(row=row, column=col)
