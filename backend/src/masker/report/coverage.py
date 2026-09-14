"""Покрытие документа и детекторов для report.json.

Перенесено из ``masker.cli`` без изменения поведения (T1.10, шаг 2): CLI
больше не должен знать, как устроен zip DOCX или части текстового слоя PDF —
это дело отчёта.
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import pymupdf
from docx import Document as open_docx

from masker.detect import DetectAgent
from masker.detect.ner import NatashaDetector
from masker.ingest.docx_ingest import (
    count_nested_tables,
    count_skipped_body_blocks,
    iter_body_blocks,
)
from masker.ingest.xlsx_ingest import PART_CELL
from masker.model import Document

WORD_TEXT_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"


def _xml_part_has_text(archive: zipfile.ZipFile, name: str) -> bool:
    root = ElementTree.fromstring(archive.read(name))
    return any((element.text or "").strip() for element in root.iter(WORD_TEXT_TAG))


def docx_coverage(source: Path, document: Document) -> dict[str, Any]:
    """Что из DOCX реально обработано — раздел ``document_coverage`` отчёта."""
    source_docx = open_docx(str(source))
    nonempty_blocks = [
        (locator, paragraph.text)
        for locator, paragraph in iter_body_blocks(source_docx)
        if paragraph.text.strip()
    ]
    body_paragraphs = [text for locator, text in nonempty_blocks if locator[0] == "body"]
    table_paragraphs = [text for locator, text in nonempty_blocks if locator[0] == "table"]
    nested_tables = count_nested_tables(source_docx)
    with zipfile.ZipFile(source) as archive:
        names = archive.namelist()
        header_names = [
            name for name in names if name.startswith("word/header") and name.endswith(".xml")
        ]
        footer_names = [
            name for name in names if name.startswith("word/footer") and name.endswith(".xml")
        ]
        header_text_parts = sum(_xml_part_has_text(archive, name) for name in header_names)
        footer_text_parts = sum(_xml_part_has_text(archive, name) for name in footer_names)
        footnotes_name = "word/footnotes.xml"
        has_footnotes_part = footnotes_name in names
        footnotes_have_text = has_footnotes_part and _xml_part_has_text(archive, footnotes_name)
    return {
        "safe_to_export": False,
        "body": {
            "processed": True,
            "nonempty_paragraphs": len(body_paragraphs),
            "skipped_blocks": count_skipped_body_blocks(source_docx),
        },
        "tables": {
            "processed": True,
            "count": len(source_docx.tables),
            "nonempty_paragraphs": len(table_paragraphs),
            "nested_count": nested_tables,
        },
        "headers": {
            "processed": False,
            "parts": len(header_names),
            "parts_with_text": header_text_parts,
        },
        "footers": {
            "processed": False,
            "parts": len(footer_names),
            "parts_with_text": footer_text_parts,
        },
        "footnotes": {
            "processed": False,
            "part_present": has_footnotes_part,
            "has_text": footnotes_have_text,
        },
        "metadata": {
            "processed": False,
            "present_fields": sorted(document.meta),
        },
    }


def pdf_coverage(source: Path, document: Document) -> dict[str, Any]:
    """Что из PDF реально обработано — раздел ``document_coverage`` отчёта."""
    pdf = pymupdf.open(str(source))  # type: ignore[no-untyped-call]
    page_count = len(pdf)
    pdf.close()  # type: ignore[no-untyped-call]
    return {
        "safe_to_export": False,
        "pages": {
            "processed": True,
            "count": page_count,
            "segment_count": len(document.segments),
        },
        "images": {
            "processed": False,
            "note": "Страницы-сканы (без текстового слоя) пропускаются (T2.3).",
        },
        "metadata": {
            "processed": False,
            "present_fields": sorted(document.meta),
        },
    }


def image_coverage(source: Path, document: Document) -> dict[str, Any]:
    """Что из картинки реально обработано — раздел ``document_coverage`` отчёта.

    Одна страница по определению; сегменты приходят из OCR-ветки
    (``origin="ocr"``). Отличается от ``pdf_coverage`` тем, что заголовок
    указывает на исходный формат картинки — читать отчёт человеку легче.
    """
    ocr_segments = sum(1 for segment in document.segments if segment.origin == "ocr")
    return {
        "safe_to_export": False,
        "image": {
            "processed": True,
            "source": document.meta.get("image_source", source.name),
            "format": document.meta.get("image_format", ""),
            "width": document.meta.get("image_width", ""),
            "height": document.meta.get("image_height", ""),
            "dpi_x": document.meta.get("image_dpi_x", ""),
            "dpi_y": document.meta.get("image_dpi_y", ""),
            "orientation": document.meta.get("image_orientation", "1"),
            "segment_count": len(document.segments),
            "ocr_segment_count": ocr_segments,
        },
        "metadata": {
            "processed": False,
            "present_fields": sorted(document.meta),
        },
    }


def xlsx_coverage(source: Path, document: Document) -> dict[str, Any]:
    """Что из XLSX реально обработано — раздел ``document_coverage`` отчёта.

    Один сегмент XLSX соответствует непустой ячейке, поэтому число сегментов
    одновременно показывает число проверенных отображаемых значений. Формулы
    учитываются отдельно: ingest читает только их кэш, а рендер обезвреживает
    зависимости от замаскированных ячеек.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(str(source), data_only=False, read_only=True)
    try:
        formula_count = sum(
            cell.data_type == "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
        return {
            "safe_to_export": False,
            "sheets": {"processed": True, "count": len(workbook.sheetnames)},
            "cells": {
                "processed": True,
                "segment_count": len(document.segments),
                "anchor_part": PART_CELL,
            },
            "formulas": {
                "processed": True,
                "count": formula_count,
                "note": "Зависимые от маски ячейки заменяются заглушкой.",
            },
            "metadata": {"processed": False, "present_fields": sorted(document.meta)},
        }
    finally:
        workbook.close()


def detection_coverage(
    selected_types: frozenset[str],
    detector: DetectAgent,
    extra_covered_types: frozenset[str] = frozenset(),
) -> dict[str, list[str]]:
    """Какие из запрошенных типов реально покрыты активными детекторами.

    ``extra_covered_types`` — типы, детектируемые вне ``detector`` (например,
    ``{"signature"}`` когда подключён отдельный детектор подписей).
    """
    active_types = {entity_type for item in detector.detectors for entity_type in item.types}
    available_types = active_types | NatashaDetector.types | extra_covered_types
    return {
        "requested_types": sorted(str(entity_type) for entity_type in selected_types),
        "active_detector_types": sorted(str(entity_type) for entity_type in active_types),
        "requested_without_detector": sorted(
            str(entity_type) for entity_type in selected_types - available_types
        ),
    }
