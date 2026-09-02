"""Контрактные тесты ingest DOCX: тело, верхнеуровневые таблицы и core properties."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document as open_docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import _BaseHeaderFooter
from docx.text.paragraph import Paragraph

from masker.detect.rules import detect_by_rules
from masker.ingest.docx_ingest import (
    count_nested_tables,
    count_skipped_body_blocks,
    ingest_docx,
    iter_runs,
    resolve_anchor,
)
from masker.model import EntityType

ROOT = next(
    parent for parent in Path(__file__).resolve().parents if (parent / "pyproject.toml").is_file()
)
FIXTURES = ROOT / "fixtures" / "labeled"
CONTRACT_01 = FIXTURES / "contract_01.docx"
CONTRACT_02 = FIXTURES / "contract_02_hard.docx"
CONTRACT_05 = FIXTURES / "contract_05_tables.docx"
CONTRACTS = [
    (CONTRACT_01, 26),
    (CONTRACT_02, 6),
]


def append_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    """Добавить внешнюю гиперссылку, как она будет сохранена в DOCX."""
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def table_shape(path: Path) -> list[tuple[int, list[int], int]]:
    """Форма таблиц по XML, включая число физических `w:tc` в строках."""
    source = open_docx(path)
    result: list[tuple[int, list[int], int]] = []
    for table in source.tables:
        grid = table._tbl.tblGrid
        result.append(
            (
                len(table._tbl.tr_lst),
                [len(row.tc_lst) for row in table._tbl.tr_lst],
                len(grid.gridCol_lst) if grid is not None else 0,
            )
        )
    return result


@pytest.mark.parametrize(("path", "expected_count"), CONTRACTS)
def test_segment_count_equals_nonempty_processed_paragraphs(
    path: Path, expected_count: int
) -> None:
    document = ingest_docx(path)

    assert len(document.segments) == expected_count


def test_table_paragraphs_become_segments_in_reading_order() -> None:
    document = ingest_docx(CONTRACT_01)

    assert len(document.segments) == 26
    assert [segment.order for segment in document.segments] == list(range(26))
    assert [segment.anchor.locator for segment in document.segments[14:23]] == [
        ("table", 0, 0, 0, 0),
        ("table", 0, 0, 1, 0),
        ("table", 0, 0, 2, 0),
        ("table", 0, 1, 0, 0),
        ("table", 0, 1, 1, 0),
        ("table", 0, 1, 2, 0),
        ("table", 0, 2, 0, 0),
        ("table", 0, 2, 1, 0),
        ("table", 0, 2, 2, 0),
    ]
    assert document.segments[13].text == "3. Спецификация"
    assert document.segments[17].text == "Резистор МЛТ-0,25"
    assert document.segments[23].text == "4. Подписи"


def test_order_is_dense_from_zero() -> None:
    document = ingest_docx(CONTRACTS[0][0])

    assert [segment.order for segment in document.segments] == list(range(len(document.segments)))


def test_body_anchors_did_not_change() -> None:
    source = open_docx(CONTRACT_01)
    document = ingest_docx(CONTRACT_01)

    body_locators = [
        segment.anchor.locator
        for segment in document.segments
        if segment.anchor.locator[0] == "body"
    ]
    assert body_locators == [
        ("body", 0),
        ("body", 1),
        ("body", 2),
        ("body", 3),
        ("body", 4),
        ("body", 5),
        ("body", 6),
        ("body", 7),
        ("body", 8),
        ("body", 9),
        ("body", 10),
        ("body", 11),
        ("body", 12),
        ("body", 13),
        ("body", 14),
        ("body", 15),
        ("body", 16),
    ]

    for segment in document.segments:
        locator = segment.anchor.locator
        if locator[0] != "body":
            continue
        para_idx = locator[1]
        assert isinstance(para_idx, int)
        assert segment.anchor.fmt == "docx"
        assert segment.anchor.label == f"абзац {para_idx + 1}"
        assert source.paragraphs[para_idx].text == segment.text


def test_segment_text_is_charwise_equal_to_paragraph_text() -> None:
    source = open_docx(CONTRACT_01)
    document = ingest_docx(CONTRACT_01)

    for segment in document.segments:
        paragraph = resolve_anchor(source, segment.anchor.locator)
        assert paragraph is not None
        assert segment.text == paragraph.text
        assert "".join(run.text for run in iter_runs(paragraph)) == segment.text


def test_table_text_is_ingested() -> None:
    document = ingest_docx(CONTRACT_01)

    assert any("Резистор МЛТ-0,25" in segment.text for segment in document.segments)
    assert "Резистор МЛТ-0,25" in document.text()


def test_meta_from_core_properties() -> None:
    document = ingest_docx(CONTRACT_01)

    assert document.meta["author"] == "Петрова Мария Сергеевна"
    assert document.meta["title"]
    assert "last_modified_by" not in document.meta
    assert "subject" not in document.meta
    assert "category" not in document.meta
    assert all(isinstance(value, str) for value in document.meta.values())


def test_document_text_joins_segments_with_newline() -> None:
    document = ingest_docx(CONTRACT_01)

    assert document.text() == "\n".join(segment.text for segment in document.segments)
    assert "ИНН 3662103003" in document.text()
    assert "Резистор МЛТ-0,25" in document.text()


@pytest.mark.parametrize("source_fixture", [CONTRACT_01, CONTRACT_05])
def test_ingest_does_not_modify_source_file(tmp_path: Path, source_fixture: Path) -> None:
    path = tmp_path / "contract.docx"
    path.write_bytes(source_fixture.read_bytes())
    before = hashlib.sha256(path.read_bytes()).digest()

    ingest_docx(path)

    assert hashlib.sha256(path.read_bytes()).digest() == before


def test_ingest_never_reads_headers_or_footers() -> None:
    with patch.object(
        _BaseHeaderFooter,
        "_get_or_add_definition",
        side_effect=AssertionError("ingest не должен читать колонтитул"),
    ):
        ingest_docx(CONTRACT_01)


@pytest.mark.parametrize("path", [CONTRACT_01, CONTRACT_05])
def test_ingest_is_deterministic(path: Path) -> None:
    expected = [(s.order, s.anchor.locator, s.text) for s in ingest_docx(path).segments]
    repeated = [(s.order, s.anchor.locator, s.text) for s in ingest_docx(path).segments]
    script = (
        "import json, sys; "
        f"sys.path.insert(0, {str(ROOT / 'src')!r}); "
        "from masker.ingest.docx_ingest import ingest_docx; "
        f"d = ingest_docx({str(path)!r}); "
        "print(json.dumps([(s.order, s.anchor.locator, s.text) for s in d.segments], "
        "ensure_ascii=False))"
    )
    result = subprocess.run(
        [sys.executable, "-c", script],
        check=True,
        capture_output=True,
        env={**os.environ, "PYTHONHASHSEED": "1"},
        text=True,
    )

    assert repeated == expected
    assert json.loads(result.stdout) == [
        [order, list(locator), text] for order, locator, text in expected
    ]


def test_runs_are_glued_without_separator(tmp_path: Path) -> None:
    path = tmp_path / "split-runs.docx"
    source = open_docx()
    paragraph = source.add_paragraph()
    for text in ("ИНН ", "36621", "03003"):
        paragraph.add_run(text)
    source.save(path)

    segment = ingest_docx(path).segments[0]
    entities = detect_by_rules([segment])

    assert segment.text == "ИНН 3662103003"
    assert any(entity.type == EntityType.INN for entity in entities)


def test_hyperlink_text_is_ingested_and_reachable_by_runs(tmp_path: Path) -> None:
    path = tmp_path / "hyperlink.docx"
    source = open_docx()
    paragraph = source.add_paragraph("Почта: ")
    append_hyperlink(paragraph, "info@triema.example", "mailto:info@triema.example")
    paragraph.add_run(" — пишите.")
    source.save(path)

    document = ingest_docx(path)
    reloaded_paragraph = open_docx(path).paragraphs[0]
    runs = iter_runs(reloaded_paragraph)

    assert "info@triema.example" in document.segments[0].text
    assert sum(len(run.text) for run in runs) == len(document.segments[0].text)
    assert len(runs) > len(reloaded_paragraph.runs)


def test_empty_paragraphs_are_skipped_but_indexes_stay_physical(tmp_path: Path) -> None:
    path = tmp_path / "empty-paragraphs.docx"
    source = open_docx()
    source.add_paragraph("")
    source.add_paragraph("Текст A")
    source.add_paragraph("   ")
    source.add_paragraph("Текст B")
    source.save(path)

    document = ingest_docx(path)

    assert [segment.order for segment in document.segments] == [0, 1]
    assert [segment.anchor.locator for segment in document.segments] == [("body", 1), ("body", 3)]


def test_table_between_paragraphs_keeps_physical_indexes(tmp_path: Path) -> None:
    path = tmp_path / "table-between-paragraphs.docx"
    source = open_docx()
    source.add_paragraph("Текст A")
    source.add_table(rows=1, cols=1).cell(0, 0).text = "ИНН 3662103003"
    source.add_paragraph("Текст B")
    source.save(path)

    document = ingest_docx(path)

    assert [segment.text for segment in document.segments] == [
        "Текст A",
        "ИНН 3662103003",
        "Текст B",
    ]
    assert [segment.anchor.locator for segment in document.segments] == [
        ("body", 0),
        ("table", 0, 0, 0, 0),
        ("body", 1),
    ]


def test_sdt_block_is_skipped_and_counted(tmp_path: Path) -> None:
    path = tmp_path / "sdt-block.docx"
    source = open_docx()
    source.add_paragraph("Текст A")
    source.add_paragraph("Текст B")
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Текст внутри SDT"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    sdt.append(content)
    source._element.body.insert(1, sdt)
    source.save(path)

    document = ingest_docx(path)

    assert [segment.text for segment in document.segments] == ["Текст A", "Текст B"]
    assert [segment.anchor.locator for segment in document.segments] == [("body", 0), ("body", 1)]
    assert count_skipped_body_blocks(open_docx(path)) == 1


def test_resolve_anchor_round_trip() -> None:
    source = open_docx(CONTRACT_01)
    document = ingest_docx(CONTRACT_01)

    for segment in document.segments:
        paragraph = resolve_anchor(source, segment.anchor.locator)
        assert paragraph is not None
        assert paragraph.text == segment.text


def test_horizontally_merged_cell_yields_single_segment() -> None:
    document = ingest_docx(CONTRACT_05)
    matches = [
        segment
        for segment in document.segments
        if segment.text == "Объединённая ячейка: ИНН 3662103003"
    ]

    assert len(matches) == 1
    assert matches[0].anchor.locator == ("table", 0, 0, 0, 0)


def test_vertically_merged_cell_is_not_duplicated() -> None:
    document = ingest_docx(CONTRACT_05)

    assert [segment.text for segment in document.segments].count("Вертикальное объединение") == 1


def test_nested_table_text_is_not_ingested_but_counted() -> None:
    source = open_docx(CONTRACT_05)
    document = ingest_docx(CONTRACT_05)

    assert "СНИЛС 112-233-445 95" not in document.text()
    assert count_nested_tables(source) == 1


def test_ingest_preserves_table_shape(tmp_path: Path) -> None:
    path = tmp_path / "contract_05_tables.docx"
    path.write_bytes(CONTRACT_05.read_bytes())
    before_shape = table_shape(path)
    before_hash = hashlib.sha256(path.read_bytes()).digest()

    ingest_docx(path)

    assert table_shape(path) == before_shape
    assert hashlib.sha256(path.read_bytes()).digest() == before_hash


def test_document_without_core_properties_gives_empty_meta(tmp_path: Path) -> None:
    path = tmp_path / "empty-document.docx"
    source = open_docx()
    properties = source.core_properties
    for name in (
        "author",
        "last_modified_by",
        "title",
        "subject",
        "comments",
        "category",
        "keywords",
    ):
        setattr(properties, name, "")
    source.save(path)

    document = ingest_docx(path)

    assert document.meta == {}
    assert document.segments == []
    assert document.text() == ""
