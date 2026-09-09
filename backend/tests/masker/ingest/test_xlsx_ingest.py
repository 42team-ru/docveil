"""Контрактные тесты ingest XLSX: ячейки, формулы, core properties, якоря."""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl

from masker.detect.agent import DetectAgent
from masker.ingest.xlsx_ingest import ingest_xlsx, resolve_anchor
from masker.model import EntityType

ROOT = next(
    parent for parent in Path(__file__).resolve().parents if (parent / "pyproject.toml").is_file()
)
FIXTURES = ROOT / "fixtures" / "labeled"
ORDER_01 = FIXTURES / "order_01.xlsx"


def test_fixture_is_ingested_with_expected_critical_entities() -> None:
    """Сквозной путь ingest → detect на реальной фикстуре корпуса (М7, критерий 1)."""
    document = ingest_xlsx(ORDER_01)
    entities = DetectAgent().detect(document).entities

    found_by_type = {e.type: e.text for e in entities}
    for etype, expected_text in (
        (EntityType.INN, "3662103003"),
        (EntityType.OGRN, "1023601546902"),
        (EntityType.SNILS, "112-233-445 95"),
        (EntityType.BANK_ACCOUNT, "40702810100000000002"),
        (EntityType.PASSPORT, "45 04 № 123456"),
    ):
        assert found_by_type.get(etype) == expected_text


def test_anchor_locator_is_cell_sheet_row_column() -> None:
    document = ingest_xlsx(ORDER_01)

    inn_segment = next(s for s in document.segments if s.text == "3662103003")

    assert inn_segment.anchor.fmt == "xlsx"
    assert inn_segment.anchor.locator == ("cell", "Реестр", 3, 2)
    assert "Реестр" in inn_segment.anchor.label
    assert "B3" in inn_segment.anchor.label or "3" in inn_segment.anchor.label


def test_order_is_dense_from_zero() -> None:
    document = ingest_xlsx(ORDER_01)

    assert [s.order for s in document.segments] == list(range(len(document.segments)))


def test_empty_cells_are_skipped() -> None:
    document = ingest_xlsx(ORDER_01)

    assert all(segment.text.strip() for segment in document.segments)


def test_second_sheet_cells_are_ingested() -> None:
    document = ingest_xlsx(ORDER_01)

    sheets = {segment.anchor.locator[1] for segment in document.segments}
    assert "Свод" in sheets


def test_resolve_anchor_round_trip() -> None:
    workbook = openpyxl.load_workbook(ORDER_01)
    document = ingest_xlsx(ORDER_01)

    for segment in document.segments:
        cell = resolve_anchor(workbook, segment.anchor.locator)
        assert cell is not None
        assert str(cell.value) == segment.text or (
            isinstance(cell.value, (int, float)) and str(int(cell.value)) == segment.text
        )


def test_resolve_anchor_rejects_unknown_locator_shapes() -> None:
    workbook = openpyxl.load_workbook(ORDER_01)

    assert resolve_anchor(workbook, ("cell", "НетТакогоЛиста", 1, 1)) is None
    assert resolve_anchor(workbook, ("cell", "Реестр", 0, 1)) is None
    assert resolve_anchor(workbook, ("body", 1)) is None


def test_ingest_does_not_modify_source_file(tmp_path: Path) -> None:
    path = tmp_path / "order.xlsx"
    path.write_bytes(ORDER_01.read_bytes())
    before = hashlib.sha256(path.read_bytes()).digest()

    ingest_xlsx(path)

    assert hashlib.sha256(path.read_bytes()).digest() == before


def test_meta_from_workbook_properties(tmp_path: Path) -> None:
    path = tmp_path / "with-author.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Текст"
    workbook.properties.creator = "Петрова Мария Сергеевна"
    workbook.properties.title = "Смета"
    workbook.save(path)

    document = ingest_xlsx(path)

    assert document.meta["author"] == "Петрова Мария Сергеевна"
    assert document.meta["title"] == "Смета"
    assert "last_modified_by" not in document.meta


def test_document_without_properties_gives_minimal_meta(tmp_path: Path) -> None:
    """`creator` пустым не бывает: `openpyxl` подставляет литеральное
    `"openpyxl"`, если элемент `dc:creator` в `docProps/core.xml`
    отсутствует (поведение библиотеки, не наша логика, — проверено
    round-trip'ом сохранения книги без явного автора). Это не утечка: это
    маркер библиотеки, а не значение, которое когда-либо ввёл человек.
    Остальные атрибуты действительно исчезают из `meta`.
    """
    path = tmp_path / "empty-props.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Текст"
    for attr in (
        "creator",
        "lastModifiedBy",
        "title",
        "subject",
        "description",
        "category",
        "keywords",
    ):
        setattr(workbook.properties, attr, None)
    workbook.save(path)

    document = ingest_xlsx(path)

    assert document.meta == {"author": "openpyxl"}


def test_formula_cell_uses_cached_value_not_formula_code(tmp_path: Path) -> None:
    """Сегмент формулы — это её результат, а не код: `=B2*1.2` не должно
    попасть в текст, который сканирует детектор (см. докстринг модуля)."""
    path = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = 100
    sheet["B1"] = "=A1*1.2"
    workbook.save(path)
    _bake_cached_value(path, "B1", "120")

    document = ingest_xlsx(path)

    b1_segment = next(s for s in document.segments if s.anchor.locator == ("cell", "Sheet", 1, 2))
    assert b1_segment.text == "120"
    assert "=" not in b1_segment.text


def test_formula_cell_without_cache_yields_no_segment(tmp_path: Path) -> None:
    """Формула, ни разу не открытая в офисном приложении, не даёт кэша —
    сканировать нечего, и ingest честно не выдумывает сегмент (см. докстринг
    модуля про то, что защита от такой ячейки живёт в render, а не здесь)."""
    path = tmp_path / "no-cache.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = 100
    sheet["B1"] = "=A1*1.2"
    workbook.save(path)

    document = ingest_xlsx(path)

    locators = [s.anchor.locator for s in document.segments]
    assert ("cell", "Sheet", 1, 2) not in locators
    assert ("cell", "Sheet", 1, 1) in locators


def test_date_cell_is_formatted_as_russian_numeric_date(tmp_path: Path) -> None:
    import datetime

    path = tmp_path / "dates.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = datetime.date(2026, 1, 15)
    workbook.save(path)

    document = ingest_xlsx(path)

    assert document.segments[0].text == "15.01.2026"


def _bake_cached_value(path: Path, coordinate: str, cached: str) -> None:
    """Вписать в готовый xlsx кэш вычисления формулы напрямую в XML.

    `openpyxl` не пишет кэш формул сам (см. докстринг `render/xlsx_redact.py`
    про побочную зачистку кэша), поэтому тест, которому кэш нужен как
    входные данные, обязан положить его туда в обход `openpyxl` — это
    единственное разрешённое место такого приёма, сам продуктовый код
    XML руками не трогает нигде.
    """
    import re
    import zipfile

    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        contents = {name: archive.read(name) for name in names}

    xml = contents["xl/worksheets/sheet1.xml"].decode("utf-8")
    pattern = re.compile(rf'(<c r="{coordinate}"[^>]*>.*?<f>[^<]*</f>)(<v>[^<]*</v>)?(</c>)')
    xml = pattern.sub(rf"\g<1><v>{cached}</v>\g<3>", xml, count=1)
    contents["xl/worksheets/sheet1.xml"] = xml.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            archive.writestr(name, contents[name])
