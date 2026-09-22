"""Тесты docx_preview: подсветка найденных сущностей в копии DOCX (не маска)."""

from __future__ import annotations

import pathlib

from docx import Document as open_docx
from docx.oxml.ns import qn

from masker.ingest.docx_ingest import ingest_docx
from masker.model import Document, Entity, EntityType, Source
from masker.render.docx_preview import render_docx_preview

_INN = "3662103003"


def _make_docx(tmp_path: pathlib.Path, text: str = f"ИНН {_INN}") -> pathlib.Path:
    path = tmp_path / "source.docx"
    doc = open_docx()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _entity_for(document: Document, text: str, etype: EntityType) -> Entity:
    seg = next(s for s in document.segments if text in s.text)
    start = seg.text.index(text)
    return Entity(
        type=etype,
        text=text,
        segment_order=seg.order,
        start=start,
        end=start + len(text),
        source=Source.RULE,
        confidence=1.0,
        normalized=text,
    )


def _shading_run(doc, needle: str):
    return next(run for paragraph in doc.paragraphs for run in paragraph.runs if needle in run.text)


def test_default_highlight_uses_configured_amber(tmp_path: pathlib.Path) -> None:
    src = _make_docx(tmp_path)
    dest = tmp_path / "preview.docx"
    document = ingest_docx(src)
    entity = _entity_for(document, _INN, EntityType.INN)

    render_docx_preview(src, dest, document, [entity])

    doc = open_docx(str(dest))
    run = _shading_run(doc, _INN)
    shd = run._r.find(f"{qn('w:rPr')}/{qn('w:shd')}")
    assert shd is not None
    assert shd.get(qn("w:fill")).upper() == "FFDE66"


def test_accepts_custom_background_and_none(tmp_path: pathlib.Path) -> None:
    src = _make_docx(tmp_path)
    document = ingest_docx(src)
    entity = _entity_for(document, _INN, EntityType.INN)

    colored = tmp_path / "colored.docx"
    render_docx_preview(src, colored, document, [entity], highlight_background="#12ab34")
    colored_run = _shading_run(open_docx(str(colored)), _INN)
    shd = colored_run._r.find(f"{qn('w:rPr')}/{qn('w:shd')}")
    assert shd is not None
    assert shd.get(qn("w:fill")) == "12AB34"

    without_background = tmp_path / "without-background.docx"
    render_docx_preview(src, without_background, document, [entity], highlight_background=None)
    plain_run = _shading_run(open_docx(str(without_background)), _INN)
    assert plain_run._r.find(f"{qn('w:rPr')}/{qn('w:shd')}") is None


def test_original_text_is_not_removed(tmp_path: pathlib.Path) -> None:
    """Предпросмотр только подсвечивает — в отличие от `docx_redact`, текст не заменяется."""
    src = _make_docx(tmp_path)
    dest = tmp_path / "preview.docx"
    document = ingest_docx(src)
    entity = _entity_for(document, _INN, EntityType.INN)

    render_docx_preview(src, dest, document, [entity])

    doc = open_docx(str(dest))
    assert _INN in " ".join(p.text for p in doc.paragraphs)
