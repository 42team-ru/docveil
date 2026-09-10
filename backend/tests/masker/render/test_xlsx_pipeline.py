"""Сквозной путь ingest → detect → plan → render на фикстуре М7 корпуса.

Критерий приёмки 1 плана М7: `fixtures/labeled/order_01.xlsx` разобрана,
сущности найдены, оба варианта (marker/blackbox) собраны, и ни один
критичный реквизит не остаётся в выходном файле — ни в тексте ячеек, ни в
формуле (`Свод!B1 = 'Реестр'!B{N}` буквально ссылается на замаскированную
ячейку с ИНН — ровно риск, описанный в AGENTS.md), ни в кэше вычислений.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import openpyxl

from masker.detect.agent import DetectAgent
from masker.ingest.xlsx_ingest import ingest_xlsx
from masker.mask.agent import PlanAgent
from masker.model import CRITICAL_TYPES
from masker.render.xlsx_redact import render_xlsx_redacted

ROOT = next(
    parent for parent in Path(__file__).resolve().parents if (parent / "pyproject.toml").is_file()
)
FIXTURE = ROOT / "fixtures" / "labeled" / "order_01.xlsx"
LABELS = ROOT / "fixtures" / "labeled" / "order_01.labels.json"


def _all_relevant_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        names = [
            name
            for name in archive.namelist()
            if name.startswith("xl/worksheets/")
            or name in ("xl/sharedStrings.xml", "xl/calcChain.xml")
        ]
        return "\n".join(archive.read(name).decode("utf-8", errors="ignore") for name in names)


def test_critical_recall_is_one_on_fixture() -> None:
    labels = json.loads(LABELS.read_text(encoding="utf-8"))
    critical_expected = {
        (entity["type"], entity["text"])
        for entity in labels["entities"]
        if entity["type"] in CRITICAL_TYPES
    }
    assert critical_expected, "фикстура должна содержать хотя бы один критичный тип"

    document = ingest_xlsx(FIXTURE)
    found = {
        (e.type, e.text)
        for e in DetectAgent().detect(document).entities
        if e.type in CRITICAL_TYPES
    }

    missing = critical_expected - found
    assert not missing, f"пропущены критичные сущности: {missing}"


def test_masked_outputs_do_not_leak_source_text_or_formulas(tmp_path: Path) -> None:
    document = ingest_xlsx(FIXTURE)
    entities = DetectAgent().detect(document).entities
    plan = PlanAgent().plan(document, entities)
    assert plan.replacements, "план не должен быть пустым на фикстуре с реквизитами"

    labels = json.loads(LABELS.read_text(encoding="utf-8"))
    secret_values = {entity["text"] for entity in labels["entities"]}

    for style in ("marker", "blackbox"):
        destination = tmp_path / f"order_01.{style}.xlsx"
        render_xlsx_redacted(FIXTURE, destination, document, plan, style=style)

        visible_text = _all_relevant_text(destination)
        for secret in secret_values:
            assert secret not in visible_text, f"{secret!r} утёк в {style}-варианте"

        # Кэш формул (data_only) не должен отдавать секрет из "Свод"/"Реестр".
        cached = openpyxl.load_workbook(destination, data_only=True)
        for sheet in cached.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        for secret in secret_values:
                            assert secret not in cell.value


def test_formula_referencing_masked_inn_cell_is_neutralized_on_fixture(tmp_path: Path) -> None:
    """Явная проверка сценария AGENTS.md на самой фикстуре корпуса, а не
    только на синтетических книгах `test_xlsx_redact.py`: `Свод!B1`
    ссылается на ячейку с ИНН листа `Реестр`."""
    document = ingest_xlsx(FIXTURE)
    entities = DetectAgent().detect(document).entities
    plan = PlanAgent().plan(document, entities)

    destination = tmp_path / "order_01.masked.xlsx"
    render_xlsx_redacted(FIXTURE, destination, document, plan, style="marker")

    result = openpyxl.load_workbook(destination)
    summary_cell = result["Свод"]["B1"]
    assert summary_cell.value != "=Реестр!B3"
    assert "3662103003" not in str(summary_cell.value)


def test_structure_is_preserved_on_fixture(tmp_path: Path) -> None:
    document = ingest_xlsx(FIXTURE)
    entities = DetectAgent().detect(document).entities
    plan = PlanAgent().plan(document, entities)

    destination = tmp_path / "order_01.masked.xlsx"
    render_xlsx_redacted(FIXTURE, destination, document, plan)

    before = openpyxl.load_workbook(FIXTURE)
    after = openpyxl.load_workbook(destination)

    assert after.sheetnames == before.sheetnames
    for name in before.sheetnames:
        assert after[name].dimensions == before[name].dimensions
        assert after[name].merged_cells.ranges == before[name].merged_cells.ranges


def test_idempotent_second_pass_finds_nothing_new(tmp_path: Path) -> None:
    document = ingest_xlsx(FIXTURE)
    entities = DetectAgent().detect(document).entities
    plan = PlanAgent().plan(document, entities)

    first = tmp_path / "first.xlsx"
    render_xlsx_redacted(FIXTURE, first, document, plan)

    document2 = ingest_xlsx(first)
    entities2 = DetectAgent().detect(document2).entities
    critical_again = [e for e in entities2 if e.type in CRITICAL_TYPES]

    assert critical_again == []
