"""Тесты `masker.eval` — метрики по корпусу (T2.2.1, шаги 2 и 3).

Пайплайн (`masker.pipeline.mask_and_validate`) и профилирование
(`masker.eval._profile_judge_metrics`) в этих тестах подменены синтетикой:
цель — проверить логику самого `eval.py` (схлопывание пробелов, разрез по
форматам, гейт на утечки и дубли маркеров), а не гонять NER/LLM/рендер ещё
раз — это уже делает `make eval` отдельно и небыстро (полный корпус,
реальные PDF/DOCX).
"""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any

import pytest
from docx import Document as WordDocument

import masker.eval as eval_module
from masker.model import (
    Anchor,
    ArtifactLayout,
    Certificate,
    CertificateCheck,
    Entity,
    EntityType,
    Leak,
    MaskGroup,
    MaskPlan,
    Replacement,
    Source,
    ValidationReport,
)
from masker.pipeline import MaskResult
from masker.run import RunFailedError

#: Заглушка `_profile_judge_metrics`: тесты этого файла не проверяют профили
#: и судью — только то, что `run()` строит из `load_corpus`/`mask_and_validate`.
_NEUTRAL_PROFILE_JUDGE_METRICS = {
    "cluster_purity": 1.0,
    "role_coverage": 1.0,
    "role_accuracy": 1.0,
    "critical_in_questions": 0.0,
    "questions_per_document": 0.0,
    "policy_questions_per_document": 0.0,
    "critical_unmasked": 0.0,
}

_EMPTY_VALIDATION = ValidationReport(
    leaked=(), residual=(), checked_artifacts=(), checked_parts=(), ok=True
)


def _replacement(entity_type: EntityType, text: str, *, ref: str = "R1") -> Replacement:
    entity = Entity(
        type=entity_type,
        text=text,
        segment_order=0,
        start=0,
        end=len(text),
        source=Source.RULE,
    )
    return Replacement(
        ref=ref,
        entity=entity,
        marker="[X]",
        group_id="G1",
        profile_id="",
        anchor=Anchor(fmt="docx", locator=("body", 0)),
    )


def _patch_corpora(
    monkeypatch: pytest.MonkeyPatch,
    *,
    main: list[tuple[Path, dict[str, Any]]] = (),
    holdout: list[tuple[Path, dict[str, Any]]] = (),
    negative: list[tuple[Path, dict[str, Any]]] = (),
) -> None:
    """Подменить `load_corpus` тремя независимыми корпусами (К2).

    `run()` зовёт `load_corpus` трижды — без аргумента для основного корпуса
    и с `FIXTURES_HOLDOUT`/`FIXTURES_NEGATIVE` для новых секций. Один
    zero-arg lambda (как было до К2) сломался бы на втором и третьем вызове
    ``TypeError``, поэтому подмена должна различать корпус по переданному пути.
    """

    def fake(fixtures: Path = eval_module.FIXTURES) -> list[tuple[Path, dict[str, Any]]]:
        if fixtures == eval_module.FIXTURES_HOLDOUT:
            return list(holdout)
        if fixtures == eval_module.FIXTURES_NEGATIVE:
            return list(negative)
        return list(main)

    monkeypatch.setattr(eval_module, "load_corpus", fake)
    monkeypatch.setattr(
        eval_module, "_profile_judge_metrics", lambda _corpus: dict(_NEUTRAL_PROFILE_JUDGE_METRICS)
    )


def _patch_common(
    monkeypatch: pytest.MonkeyPatch, corpus: list[tuple[Path, dict[str, Any]]]
) -> None:
    """Совместимость со старыми тестами: только основной корпус, holdout и
    негативный корпус пусты — секции К2 печатают «ПРОПУЩЕН» и не влияют
    на ворота (см. `test_eval_holdout_and_negative_sections_skip_when_empty`).
    """
    _patch_corpora(monkeypatch, main=corpus)


def _row(output: str, prefix: str) -> str:
    line = next((line for line in output.splitlines() if line.startswith(prefix)), None)
    assert line is not None, f"строка {prefix!r} не найдена в выводе:\n{output}"
    return line


def _patch_mask_and_validate(
    monkeypatch: pytest.MonkeyPatch, by_path: dict[str, MaskResult]
) -> None:
    """Подменить `masker.pipeline.mask_and_validate` на заранее заготовленные результаты."""

    @contextmanager
    def fake(path: Path, *, types: Any, custom_types: Any = ()) -> Iterator[MaskResult]:
        yield by_path[str(path)]

    monkeypatch.setattr("masker.pipeline.mask_and_validate", fake)


def test_eval_collapses_whitespace_in_labels(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Разметка с лишними/непохожими пробелами не должна давать FN.

    Разметка PDF-корпуса пишется в «естественной однострочной форме» (схема
    разметки, пункт 1) и не обязана совпадать с тем, как ingest режет текст
    на сегменты. Без схлопывания пробелов двойной пробел в `text` разметки
    увёл бы это значение в false negative, хотя оно найдено верно.
    """
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    # Разметка — с двойным пробелом; найденная сущность — с одним пробелом,
    # ровно тот же класс расхождения, что описан в схеме разметки PDF.
    labels = {"entities": [{"type": "org_name", "text": "ООО  «Ромашка»"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.ORG_NAME, "ООО «Ромашка»"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    eval_module.run(gate=False)
    output = capsys.readouterr().out
    row = _row(output, "org_name")
    _, _precision, recall, _f1, fn, _fp = row.split()
    assert recall == "1.000", f"ожидали recall 1.0 при схлопывании пробелов, строка: {row!r}"
    assert fn == "0", f"двойной пробел в разметке не должен давать FN, строка: {row!r}"


def test_eval_reports_pdf_format_row(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Секция «ФОРМАТЫ» обязана появиться и содержать строку `pdf`.

    До шага 2 `_profile_judge_metrics` пропускал всё, что не `.docx`
    (`if path.suffix != ".docx": continue` — Д7 плана T2.2.1), а `run()` не
    строил разреза по форматам вовсе — идеальные цифры по DOCX маскировали
    провал по PDF.
    """
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"")
    corpus = [
        (docx_path, {"entities": [{"type": "inn", "text": "1234567890"}]}),
        (pdf_path, {"entities": [{"type": "inn", "text": "1234567890"}]}),
    ]
    _patch_common(monkeypatch, corpus)
    # PDF нарочно не находит ничего — именно так выглядит Д7 сегодня: PDF
    # физически участвует в прогоне, но детекция по нему хромает.
    docx_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    pdf_result = MaskResult(
        plan=MaskPlan(replacements=(), groups=(), skipped=(), requested_types=()),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): docx_result, str(pdf_path): pdf_result})

    eval_module.run(gate=False)
    output = capsys.readouterr().out
    assert "ФОРМАТЫ" in output
    docx_row = _row(output, "docx")
    pdf_row = _row(output, "pdf")
    assert docx_row.split()[2] == "1.000", f"docx recall должен остаться 1.0: {docx_row!r}"
    assert pdf_row.split()[2] == "0.000", f"pdf recall должен показать провал: {pdf_row!r}"


def test_eval_gate_fails_on_leak(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Синтетический план с заведомой утечкой — ворота обязаны провалиться.

    До шага 3 `leaked_total` нигде не читался: `ValidateAgent` уже находил
    утечки (T1.8), но `eval.py` их не суммировал и не проверял порогом.
    """
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    leak = Leak(
        kind="raw",
        artifact="masked_black.docx",
        part="word/document.xml",
        entity_type="inn",
        value="1234567890",
        ref="R1",
        group_id="G1",
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(leak,),
            residual=(),
            checked_artifacts=("masked_black.docx",),
            checked_parts=(),
            ok=False,
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "утечка обязана провалить ворота"
    assert _row(output, "leaked_total").split()[-1] == "1"


def test_eval_gate_passes_without_leak(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Симметричный случай: без утечки ворота по `leaked_total` не падают."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert _row(output, "leaked_total").split()[-1] == "0"


def test_eval_gate_fails_on_certificate_failure(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """План М3: сертификат обезличивания провалил пункт `width_quantization` —
    ворота обязаны провалиться на `certificate_failures`, даже когда
    `leaked_total`/`layout_removed_chars` в порядке."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    certificate = Certificate(
        ok=False,
        checks=(
            CertificateCheck(name="leak_scan", ok=True, detail="утечек не найдено"),
            CertificateCheck(name="metadata_cleared", ok=True, detail="метаданные пусты"),
            CertificateCheck(
                name="width_quantization",
                ok=False,
                detail="R1 (inn, стр. 1): ширина 62.00pt не кратна 12pt",
            ),
        ),
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(),
            residual=(),
            checked_artifacts=("masked_black.docx",),
            checked_parts=(),
            ok=True,
            certificate=certificate,
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "провал сертификата обязан провалить ворота"
    assert _row(output, "certificate_failures").split()[-1] == "1"
    assert "width_quantization" in output


def test_eval_gate_passes_with_certificate_ok(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Симметричный случай: сертификат пройден целиком — ворота по
    `certificate_failures` не падают."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    certificate = Certificate(
        ok=True,
        checks=(
            CertificateCheck(name="leak_scan", ok=True, detail="утечек не найдено"),
            CertificateCheck(name="metadata_cleared", ok=True, detail="метаданные пусты"),
            CertificateCheck(name="width_quantization", ok=True, detail="кратно 12pt"),
        ),
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(),
            residual=(),
            checked_artifacts=("masked_black.docx",),
            checked_parts=(),
            ok=True,
            certificate=certificate,
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert _row(output, "certificate_failures").split()[-1] == "0"


def test_eval_gate_fails_on_layout_loss(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Прямоугольник редакции стёр текст вне своих замен (Д10, план T2.2.2,
    шаг 5) — ворота обязаны провалиться на `layout_removed_chars`."""
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(pdf_path, labels)])
    layout = ArtifactLayout(
        artifact="masked_black.pdf",
        removed_chars=46,
        inserted_chars=0,
        pages=(0,),
        first_diff="стр. 1: ожидалось …Устава, с одной…, получено …Устава…",
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(),
            residual=(),
            checked_artifacts=("masked_black.pdf",),
            checked_parts=(),
            ok=True,
            layout=(layout,),
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(pdf_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "потеря текста вне замены обязана провалить ворота"
    assert _row(output, "layout_removed_chars").split()[-1] == "46"
    assert "masked_black.pdf" in output


def test_eval_gate_passes_without_layout_loss(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Симметричный случай: без потери вёрстки ворота по
    `layout_removed_chars` не падают."""
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(pdf_path, labels)])
    layout = ArtifactLayout(
        artifact="masked_black.pdf",
        removed_chars=0,
        inserted_chars=0,
        pages=(),
        first_diff="",
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(),
            residual=(),
            checked_artifacts=("masked_black.pdf",),
            checked_parts=(),
            ok=True,
            layout=(layout,),
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(pdf_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert _row(output, "layout_removed_chars").split()[-1] == "0"


def test_eval_gate_fails_on_duplicate_marker(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Маркер, вставленный дважды на один `Replacement` (Д1), обязан провалить ворота."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "org_name", "text": "ООО «Ромашка»"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])

    artifact = tmp_path / "masked_black.docx"
    word_doc = WordDocument()
    # Один Replacement, но маркер физически вставлен в артефакт дважды —
    # ровно дефект Д1 (рендер искал строку заново и попал в неё повторно).
    word_doc.add_paragraph("[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]")
    word_doc.add_paragraph("[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]")
    word_doc.save(artifact)

    plan = MaskPlan(
        replacements=(_replacement(EntityType.ORG_NAME, "ООО «Ромашка»"),),
        groups=(
            MaskGroup(
                id="G1",
                key="org:1",
                type=EntityType.ORG_NAME,
                marker="[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]",
                profile_id="P1",
                role_label="ПОСТАВЩИК",
                number=1,
                refs=("R1",),
                sample="ООО «Ромашка»",
            ),
        ),
        skipped=(),
        requested_types=(),
    )
    result = MaskResult(plan=plan, validation=_EMPTY_VALIDATION, artifacts=(artifact,))
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "дубль маркера обязан провалить ворота"
    assert _row(output, "duplicate_markers").split()[-1] == "1"


def test_duplicate_marker_count_ignores_exact_match(tmp_path: Path) -> None:
    """Маркер, вставленный ровно по разу на `Replacement`, — не дубль."""
    artifact = tmp_path / "masked_black.docx"
    word_doc = WordDocument()
    word_doc.add_paragraph("[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]")
    word_doc.save(artifact)
    plan = MaskPlan(
        replacements=(),
        groups=(
            MaskGroup(
                id="G1",
                key="org:1",
                type=EntityType.ORG_NAME,
                marker="[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]",
                profile_id="P1",
                role_label="ПОСТАВЩИК",
                number=1,
                refs=("R1",),
                sample="ООО «Ромашка»",
            ),
        ),
        skipped=(),
        requested_types=(),
    )
    assert eval_module.duplicate_marker_count(plan, (artifact,)) == 0


def _group_with_ladder(*, refs: tuple[str, ...] = ("R1", "R2")) -> MaskGroup:
    """Группа ФИО с ролью, два вхождения — годится для обеих ступеней лестницы."""
    return MaskGroup(
        id="G1",
        key="person:1",
        type=EntityType.PERSON,
        marker="[ПОСТАВЩИК-ФИО]",
        profile_id="P1",
        role_label="ПОСТАВЩИК",
        number=1,
        refs=refs,
        sample="Иванов",
        canonical_label="[Поставщик Представитель]",
        compact_label="[Ф1]",
    )


def _degradation(
    group_id: str, shown_label: str, *, fallback_reason: str = "compact"
) -> dict[str, Any]:
    """Один элемент `render_degradations` — контракт `graph/nodes.py::render_node`."""
    return {
        "artifact": "masked_highlight.pdf",
        "role": "highlight",
        "page": 0,
        "group_id": group_id,
        "entity_type": EntityType.PERSON,
        "canonical_label": "[Поставщик Представитель]",
        "shown_label": shown_label,
        "font_size": 9.0,
        "fallback_reason": fallback_reason,
    }


def test_inconsistent_marker_count_flags_two_labels_for_one_group() -> None:
    """План М4: одна и та же группа не должна печататься двумя разными
    строками — встречная проверка к `duplicate_marker_count`. Одно вхождение
    спустилось по лестнице (``[Ф1]``), второе показало канонический маркер
    как есть (не попадает в `render_degradations` — норма, план М1)."""
    plan = MaskPlan(
        replacements=(),
        groups=(_group_with_ladder(refs=("R1", "R2")),),
        skipped=(),
        requested_types=(),
    )
    render_degradations = (_degradation("G1", "[Ф1]"),)
    assert eval_module.inconsistent_marker_count(plan, render_degradations) == 1


def test_inconsistent_marker_count_ignores_group_with_one_label() -> None:
    """Оба вхождения группы спустились до одной и той же ступени — не дефект."""
    plan = MaskPlan(
        replacements=(),
        groups=(_group_with_ladder(refs=("R1", "R2")),),
        skipped=(),
        requested_types=(),
    )
    render_degradations = (_degradation("G1", "[Ф1]"), _degradation("G1", "[Ф1]"))
    assert eval_module.inconsistent_marker_count(plan, render_degradations) == 0


def test_inconsistent_marker_count_ignores_group_without_degradations() -> None:
    """Ни одно вхождение не спустилось по лестнице — все показали канонический
    маркер как есть, `render_degradations` для группы пуст."""
    plan = MaskPlan(
        replacements=(),
        groups=(_group_with_ladder(refs=("R1", "R2")),),
        skipped=(),
        requested_types=(),
    )
    assert eval_module.inconsistent_marker_count(plan, ()) == 0


def test_eval_gate_fails_on_inconsistent_marker(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Группа, отрендеренная двумя разными строками, обязана провалить ворота."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "person", "text": "Иванов"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])

    artifact = tmp_path / "masked_highlight.docx"
    word_doc = WordDocument()
    word_doc.add_paragraph("[Поставщик Представитель]")
    word_doc.save(artifact)

    plan = MaskPlan(
        replacements=(),
        groups=(_group_with_ladder(refs=("R1", "R2")),),
        skipped=(),
        requested_types=(),
    )
    result = MaskResult(
        plan=plan,
        validation=_EMPTY_VALIDATION,
        artifacts=(artifact,),
        render_degradations=(_degradation("G1", "[Ф1]"),),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "рассогласованный маркер обязан провалить ворота"
    assert _row(output, "inconsistent_markers").split()[-1] == "1"


def test_eval_gate_fails_on_render_failure_but_keeps_measuring_the_rest(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Прогон корпуса — измерительный инструмент: аномальная сущность на
    одном документе не имеет права ослепить ворота целиком и скрыть
    метрики по остальным документам (план T2.2.1, пачка 4). При этом
    падение обязано остаться видимым — новой метрикой с адресом проблемы,
    а не тихим пропуском документа.
    """
    broken_path = tmp_path / "broken.pdf"
    broken_path.write_bytes(b"")
    healthy_path = tmp_path / "healthy.docx"
    healthy_path.write_bytes(b"")
    corpus = [
        (broken_path, {"entities": [{"type": "org_name", "text": "ГО и ЧС"}]}),
        (healthy_path, {"entities": [{"type": "inn", "text": "1234567890"}]}),
    ]
    _patch_common(monkeypatch, corpus)
    healthy_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )

    @contextmanager
    def fake(path: Path, *, types: Any, custom_types: Any = ()) -> Iterator[MaskResult]:
        if str(path) == str(broken_path):
            cause = ValueError(
                "маркер '[СТОРОНА-27-ОРГАНИЗАЦИЯ]' не помещается в прямоугольник "
                "ни при одном размере шрифта"
            )
            raise RunFailedError("thread123", "render", cause)
        yield healthy_result

    monkeypatch.setattr("masker.pipeline.mask_and_validate", fake)

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out

    assert code == 1, "падение рендера обязано провалить ворота"
    assert _row(output, "render_failures").split()[-1] == "1"
    assert broken_path.name in output, "имя упавшего документа обязано быть в отчёте"
    assert "СТОРОНА-27-ОРГАНИЗАЦИЯ" in output, "маркер, на котором упал рендер, обязан быть виден"
    # Здоровый документ не должен пострадать от падения соседнего.
    inn_row = _row(output, "inn")
    assert inn_row.split()[2] == "1.000", (
        f"метрика по healthy.docx обязана посчитаться: {inn_row!r}"
    )


def test_eval_gate_passes_without_render_failure(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Симметричный случай: без падений рендера ворота по `render_failures` не падают."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert _row(output, "render_failures").split()[-1] == "0"


def test_eval_holdout_and_negative_sections_skip_when_empty(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """К2: пустой holdout/негативный корпус не валит ворота — печатается
    явный пропуск, а не тихая нулевая метрика."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert "HOLDOUT ПРОПУЩЕН" in output
    assert "ЛОЖНЫЕ ПРОПУЩЕНЫ" in output


def test_eval_holdout_section_prints_separately_and_passes(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """К2: секция HOLDOUT печатается отдельно от основного корпуса и не
    валит ворота, когда holdout recall не хуже установленных порогов."""
    main_path = tmp_path / "main.docx"
    main_path.write_bytes(b"")
    main_labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    holdout_path = tmp_path / "holdout.docx"
    holdout_path.write_bytes(b"")
    holdout_labels = {"entities": [{"type": "inn", "text": "9876543210"}]}
    _patch_corpora(
        monkeypatch, main=[(main_path, main_labels)], holdout=[(holdout_path, holdout_labels)]
    )
    main_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    holdout_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "9876543210"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(
        monkeypatch, {str(main_path): main_result, str(holdout_path): holdout_result}
    )

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert "HOLDOUT (fixtures/holdout" in output
    assert "holdout_recall_critical     1.000  (1/1)" in output


def test_eval_holdout_gate_fails_when_critical_type_missed(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Пропуск критичного типа (ИНН) на holdout — утечка, ворота обязаны
    провалиться независимо от того, как выглядит основной корпус."""
    main_path = tmp_path / "main.docx"
    main_path.write_bytes(b"")
    main_labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    holdout_path = tmp_path / "holdout.docx"
    holdout_path.write_bytes(b"")
    holdout_labels = {"entities": [{"type": "inn", "text": "9876543210"}]}
    _patch_corpora(
        monkeypatch, main=[(main_path, main_labels)], holdout=[(holdout_path, holdout_labels)]
    )
    main_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    # Holdout ничего не находит — критичный ИНН пропущен.
    holdout_result = MaskResult(
        plan=MaskPlan(replacements=(), groups=(), skipped=(), requested_types=()),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(
        monkeypatch, {str(main_path): main_result, str(holdout_path): holdout_result}
    )

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "пропуск критичного типа на holdout обязан провалить ворота"
    assert "критичный тип не найден" in output


def test_eval_holdout_gate_fails_when_recall_other_below_threshold(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Некритичный тип, полностью пропущенный на holdout, обязан провалить
    ворота через агрегированный `holdout_recall_other`."""
    main_path = tmp_path / "main.docx"
    main_path.write_bytes(b"")
    main_labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    holdout_path = tmp_path / "holdout.docx"
    holdout_path.write_bytes(b"")
    holdout_labels = {"entities": [{"type": "org_name", "text": "ООО «Ромашка»"}]}
    _patch_corpora(
        monkeypatch, main=[(main_path, main_labels)], holdout=[(holdout_path, holdout_labels)]
    )
    main_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    holdout_result = MaskResult(
        plan=MaskPlan(replacements=(), groups=(), skipped=(), requested_types=()),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(
        monkeypatch, {str(main_path): main_result, str(holdout_path): holdout_result}
    )

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "нулевой recall по некритичному типу на holdout обязан провалить ворота"
    assert any("holdout_recall_other" in line for line in output.splitlines())


def test_eval_negative_section_passes_within_threshold(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """К2: негативный корпус без единой PII — найденное считается только как
    FP; в пределах порога ворота не падают, а секция печатается отдельно."""
    main_path = tmp_path / "main.docx"
    main_path.write_bytes(b"")
    main_labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    negative_path = tmp_path / "negative.docx"
    negative_path.write_bytes(b"")
    negative_labels: dict[str, Any] = {"entities": []}
    _patch_corpora(
        monkeypatch,
        main=[(main_path, main_labels)],
        negative=[(negative_path, negative_labels)],
    )
    main_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    # Один ложный org_name — в пределах MAX_NEGATIVE_FALSE_POSITIVES (2).
    negative_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.ORG_NAME, "ГОСТ Р 12345-2020"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(
        monkeypatch, {str(main_path): main_result, str(negative_path): negative_result}
    )

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert "ЛОЖНЫЕ (fixtures/negative" in output
    assert _row(output, "ИТОГО").split()[-1] == "1"


def test_eval_negative_section_fails_gate_above_threshold(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Ложных срабатываний на негативном корпусе больше порога — ворота
    обязаны провалиться, а не молча стерпеть рост."""
    main_path = tmp_path / "main.docx"
    main_path.write_bytes(b"")
    main_labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    negative_path = tmp_path / "negative.docx"
    negative_path.write_bytes(b"")
    negative_labels: dict[str, Any] = {"entities": []}
    _patch_corpora(
        monkeypatch,
        main=[(main_path, main_labels)],
        negative=[(negative_path, negative_labels)],
    )
    main_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    # Три ложных срабатывания — больше MAX_NEGATIVE_FALSE_POSITIVES (2).
    negative_result = MaskResult(
        plan=MaskPlan(
            replacements=(
                _replacement(EntityType.ORG_NAME, "ГОСТ Р 12345-2020", ref="R1"),
                _replacement(EntityType.DATE, "01.07.2020", ref="R2"),
                _replacement(EntityType.PERSON, "ГГ-ММ-НННН", ref="R3"),
            ),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(
        monkeypatch, {str(main_path): main_result, str(negative_path): negative_result}
    )

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "рост ложных срабатываний на негативном корпусе обязан провалить ворота"
    assert "негативный корпус" in output
    assert _row(output, "ИТОГО").split()[-1] == "3"
