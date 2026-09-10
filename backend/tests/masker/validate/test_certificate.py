"""Сертификат обезличивания (план М3) — три независимые проверки итогового файла.

``test_verify_width_quantization_fails_on_deliberately_unquantized_width`` —
ключевой тест задачи М3: доказывает, что пункт 3 (ширина не выдаёт длину
оригинала) реально нагружен, а не декоративен — сфабрикованная ширина,
пропорциональная длине текста и не объяснённая ни сеткой 12pt, ни настоящим
соседним символом на странице, обязана провалить проверку.
"""

from __future__ import annotations

import pathlib

import pymupdf
import pytest
from docx import Document as open_docx
from openpyxl import Workbook

from masker.ingest.pdf_ingest import PageChars, ingest_pdf
from masker.mask.agent import PlanAgent
from masker.model import Document, Entity, EntityType, Leak, Source
from masker.render.pdf_render import render_pdf_redacted
from masker.validate.certificate import (
    _check_leak_scan,
    _check_metadata_cleared,
    _check_width_quantization,
    build_certificate,
    verify_width_quantization,
)

_FONT = str(pathlib.Path(__file__).parent.parent.parent.parent / "src/masker/data/DejaVuSans.ttf")


def _person_entity(document: Document, name: str) -> Entity:
    seg = next(s for s in document.segments if name in s.text)
    start = seg.text.index(name)
    return Entity(
        type=EntityType.PERSON,
        text=name,
        segment_order=seg.order,
        start=start,
        end=start + len(name),
        source=Source.RULE,
        confidence=1.0,
        normalized=name.lower(),
    )


# ── verify_width_quantization: чистая функция, без открытия PDF ──────────────────


def test_verify_width_quantization_passes_grid_aligned_width() -> None:
    """Ширина, кратная сетке 12pt, проходит независимо от геометрии страницы —
    сосед вообще не нужен, когда квант применён полностью."""
    widths = [("R1", "person", 0, 0.0, 0.0, 36.0, 10.0)]
    empty_page = PageChars(text="", boxes=(), line_ids=())
    check = verify_width_quantization(widths, chars_by_page={0: empty_page})
    assert check.ok is True
    assert check.name == "width_quantization"


def test_verify_width_quantization_passes_when_capped_by_real_neighbour() -> None:
    """Ширина не кратна сетке, но на странице действительно есть символ,
    чья граница совпадает с правым краем региона, — легитимный недокрученный
    квант (план М1, правило 5: квант приносится в жертву целостности соседа)."""
    widths = [("R1", "person", 0, 0.0, 0.0, 37.3, 10.0)]
    neighbour = pymupdf.Rect(37.3, 0.0, 45.0, 10.0)
    chars = PageChars(text="X", boxes=(neighbour,), line_ids=(0,))
    check = verify_width_quantization(widths, chars_by_page={0: chars})
    assert check.ok is True


def test_verify_width_quantization_fails_on_deliberately_unquantized_width() -> None:
    """Ключевой тест плана М3: ширины, сфабрикованные строго пропорционально
    длине текста (``6.2 * len(text)``), не кратны сетке 12pt и не объяснены
    никаким символом страницы (она вообще пуста, места сколько угодно) —
    ровно тот канал утечки, для которого существует квантование М1, шаг 5.
    Проверка обязана его поймать."""
    widths = [
        ("R1", "person", 0, 0.0, 0.0, 9 * 6.2, 10.0),
        ("R2", "person", 0, 0.0, 20.0, 10 * 6.2, 30.0),
        ("R3", "person", 0, 0.0, 40.0, 11 * 6.2, 50.0),
    ]
    empty_page = PageChars(text="", boxes=(), line_ids=())
    check = verify_width_quantization(widths, chars_by_page={0: empty_page})
    assert check.ok is False
    assert "квантование" in check.detail
    assert "R1" in check.detail and "R2" in check.detail and "R3" in check.detail


def test_verify_width_quantization_missing_page_chars_treated_as_no_neighbour() -> None:
    """Если для страницы региона нет ``chars_by_page`` вовсе (не должно
    происходить в проде, но защищаемся), непокрытая нестандартная ширина не
    получает бесплатного оправдания — тоже провал."""
    widths = [("R1", "person", 5, 0.0, 0.0, 37.3, 10.0)]
    check = verify_width_quantization(widths, chars_by_page={})
    assert check.ok is False


# ── _check_width_quantization: интеграция с реальным PDF ─────────────────────────


def test_check_width_quantization_ok_on_real_quantized_document(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "surnames.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="dvu", fontfile=_FONT)
    page.insert_text((72, 100), "Попов        далее", fontname="dvu", fontsize=12)
    page.insert_text((72, 130), "Иванов        далее", fontname="dvu", fontsize=12)
    doc.save(str(path))
    doc.close()

    document = ingest_pdf(path)
    entities = [_person_entity(document, "Попов"), _person_entity(document, "Иванов")]
    plan = PlanAgent().plan(document, entities)
    dest = tmp_path / "redacted.pdf"
    render_pdf_redacted(path, dest, document, plan, style="blackbox")

    check = _check_width_quantization(plan, path, [dest])
    assert check.ok is True
    assert "проверено" in check.detail


def test_check_width_quantization_not_applicable_for_docx_source(tmp_path: pathlib.Path) -> None:
    from masker.model import MaskPlan

    plan = MaskPlan(replacements=(), groups=(), skipped=(), requested_types=())
    check = _check_width_quantization(plan, tmp_path / "source.docx", [tmp_path / "out.docx"])
    assert check.ok is True
    assert "не применимо" in check.detail


def test_check_width_quantization_fails_when_render_leaves_geometry_unquantized(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Сквозной сценарий: если бы ``render/pdf_render.py`` перестал квантовать
    (баг в будущем), сертификат обязан это заметить, а не молча согласиться —
    подменяем пересчёт геометрии на заведомо неквантованный, без изменения
    самого рендера."""
    import masker.validate.certificate as certificate_module
    from masker.model import PdfRegion

    path = tmp_path / "source.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="dvu", fontfile=_FONT)
    page.insert_text((72, 100), "Иванов", fontname="dvu", fontsize=12)
    doc.save(str(path))
    doc.close()

    document = ingest_pdf(path)
    entity = _person_entity(document, "Иванов")
    plan = PlanAgent().plan(document, [entity])
    dest = tmp_path / "redacted.pdf"
    render_pdf_redacted(path, dest, document, plan, style="blackbox")

    ref = plan.replacements[0].ref
    fake_width = len(entity.text) * 6.2  # пропорционально длине, не кратно 12pt

    def fake_geometry(_source: object, _plan: object) -> dict[str, tuple[PdfRegion, ...]]:
        return {ref: (PdfRegion(page=0, x0=0.0, y0=0.0, x1=fake_width, y1=10.0),)}

    monkeypatch.setattr(certificate_module, "_compute_erase_geometry", fake_geometry)

    certificate = build_certificate(plan, (), (), [dest], source=path)
    assert certificate.ok is False
    by_name = {check.name: check for check in certificate.checks}
    assert by_name["width_quantization"].ok is False
    # Остальные два пункта не пострадали от подмены геометрии.
    assert by_name["leak_scan"].ok is True
    assert by_name["metadata_cleared"].ok is True


# ── _check_leak_scan ───────────────────────────────────────────────────────────


def test_check_leak_scan_ok_without_leaks() -> None:
    check = _check_leak_scan((), ("word/document.xml",))
    assert check.ok is True
    assert "утечек не найдено" in check.detail


def test_check_leak_scan_fails_with_leaks() -> None:
    leak = Leak(
        kind="raw",
        artifact="masked_black.docx",
        part="word/document.xml",
        entity_type="inn",
        value="3662103003",
        ref="R1",
        group_id="G1",
    )
    check = _check_leak_scan((leak,), ("word/document.xml",))
    assert check.ok is False
    assert "1" in check.detail
    assert "inn" in check.detail


# ── _check_metadata_cleared ────────────────────────────────────────────────────


def _make_clean_docx(path: pathlib.Path) -> None:
    """docx с явно вычищенными свойствами — python-docx сам по умолчанию
    заполняет ``author``/``comments`` шаблонными значениями
    (``"python-docx"``/``"generated by python-docx"``), и это ровно то, что
    ``render/docx_redact.py`` перезаписывает пустой строкой на реальном
    рендере. «Чистый» фикстурный docx обязан воспроизводить именно
    результат рендера, а не сырой ``open_docx()``."""
    doc = open_docx()
    doc.add_paragraph("текст без реквизитов")
    props = doc.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "keywords", "comments"):
        setattr(props, attr, "")
    doc.save(str(path))


def test_check_metadata_cleared_ok_on_clean_docx(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "clean.docx"
    _make_clean_docx(path)
    check = _check_metadata_cleared([path])
    assert check.ok is True


def test_check_metadata_cleared_fails_on_dirty_docx_author(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "dirty.docx"
    doc = open_docx()
    doc.add_paragraph("текст")
    doc.core_properties.author = "Иванов Иван Иванович"
    doc.save(str(path))
    check = _check_metadata_cleared([path])
    assert check.ok is False
    assert "author" in check.detail


def test_check_metadata_cleared_ok_on_clean_pdf(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "clean.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.set_metadata({})
    doc.del_xml_metadata()
    doc.save(str(path))
    doc.close()
    check = _check_metadata_cleared([path])
    assert check.ok is True


def test_check_metadata_cleared_fails_on_dirty_pdf_author(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "dirty.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.set_metadata({"author": "Иванов Иван Иванович"})
    doc.save(str(path))
    doc.close()
    check = _check_metadata_cleared([path])
    assert check.ok is False
    assert "author" in check.detail


def test_check_metadata_cleared_rejects_unknown_format(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "artifact.txt"
    path.write_bytes(b"")
    with pytest.raises(ValueError, match="txt"):
        _check_metadata_cleared([path])


def test_check_metadata_cleared_accepts_xlsx_tool_generated_creator(
    tmp_path: pathlib.Path,
) -> None:
    """openpyxl сам записывает ``creator=openpyxl`` после очистки автора.

    Это техническая метка библиотеки, не исходная персональная метаинформация;
    сертификат не должен превращать каждый честный XLSX-рендер в провал.
    """
    path = tmp_path / "clean.xlsx"
    workbook = Workbook()
    workbook.properties.creator = "openpyxl"
    workbook.save(path)
    workbook.close()

    check = _check_metadata_cleared([path])

    assert check.ok is True


def test_check_metadata_cleared_fails_on_dirty_xlsx_creator(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "dirty.xlsx"
    workbook = Workbook()
    workbook.properties.creator = "Иванов Иван Иванович"
    workbook.save(path)
    workbook.close()

    check = _check_metadata_cleared([path])

    assert check.ok is False
    assert "creator" in check.detail


# ── build_certificate: конъюнкция независимых пунктов ─────────────────────────


def test_build_certificate_ok_only_if_all_checks_pass(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "clean.docx"
    _make_clean_docx(path)
    from masker.model import MaskPlan

    plan = MaskPlan(replacements=(), groups=(), skipped=(), requested_types=())
    certificate = build_certificate(plan, (), (), [path], source=None)
    assert certificate.ok is True
    # feat-image-ingest добавил 4-й пункт — `image_metadata_stripped`.
    # На DOCX-артефактах он всегда «не применимо».
    assert len(certificate.checks) == 4
    names = {check.name for check in certificate.checks}
    assert names == {
        "leak_scan",
        "metadata_cleared",
        "width_quantization",
        "image_metadata_stripped",
    }


def test_build_certificate_fails_if_leak_scan_fails(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "clean.docx"
    _make_clean_docx(path)
    from masker.model import MaskPlan

    plan = MaskPlan(replacements=(), groups=(), skipped=(), requested_types=())
    leak = Leak(
        kind="raw",
        artifact="clean.docx",
        part="word/document.xml",
        entity_type="inn",
        value="3662103003",
    )
    certificate = build_certificate(plan, (leak,), (), [path], source=None)
    assert certificate.ok is False
