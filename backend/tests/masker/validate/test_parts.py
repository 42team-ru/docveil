"""Тесты validate/parts.py: разбор частей контейнера (T1.8, шаг 8)."""

from __future__ import annotations

import pathlib
import zipfile

import pymupdf
from docx import Document as open_docx

from masker.validate.parts import docx_parts, pdf_parts

_AUTHOR = "Тест Автор"
_INN = "3662103003"


def _make_docx(tmp_path: pathlib.Path) -> pathlib.Path:
    path = tmp_path / "source.docx"
    doc = open_docx()
    doc.core_properties.author = _AUTHOR
    doc.add_paragraph(f"ИНН {_INN}")
    doc.save(str(path))
    return path


def test_docx_parts_cover_all_zip_entries(tmp_path: pathlib.Path) -> None:
    path = _make_docx(tmp_path)
    parts = docx_parts(path)
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
    assert {part.name for part in parts} == names


def test_docx_parts_are_sorted(tmp_path: pathlib.Path) -> None:
    path = _make_docx(tmp_path)
    names = [part.name for part in docx_parts(path)]
    assert names == sorted(names)
    assert len(names) > 1  # тест не должен пройти случайно на пустом списке


def test_core_xml_author_is_visible_in_raw(tmp_path: pathlib.Path) -> None:
    path = _make_docx(tmp_path)
    parts = {part.name: part for part in docx_parts(path)}
    core = parts["docProps/core.xml"]
    assert _AUTHOR.encode("utf-8") in core.raw
    assert _AUTHOR in core.text


def test_document_text_joins_split_runs(tmp_path: pathlib.Path) -> None:
    """Доказательство того, что одного побайтового поиска мало (T1.8):
    значение, разбитое на два run'а, целиком присутствует в `DocPart.text`
    (`ingest_docx` склеивает run'ы в текст абзаца), но не встречается в
    `raw` — XML хранит два отдельных `<w:t>` без склейки, между ними лежит
    закрывающий/открывающий тег run'а.

    Тест обязан падать на реализации, которая вообще не строит `text`
    (только побайтовый поиск, `DocPart.text` всегда пустая строка) —
    проверено вручную (см. отчёт кодера): `test_core_xml_author_is_visible_in_raw`
    и этот тест оба падают, если `docx_parts` не заполняет `text`.
    """
    path = tmp_path / "multirun.docx"
    doc = open_docx()
    para = doc.add_paragraph()
    half = len(_INN) // 2
    para.add_run(f"ИНН {_INN[:half]}")
    para.add_run(_INN[half:])
    doc.save(str(path))

    parts = {part.name: part for part in docx_parts(path)}
    document_part = parts["word/document.xml"]
    assert _INN in document_part.text
    assert _INN.encode("utf-8") not in document_part.raw


def test_pdf_parts_have_page_text_and_metadata(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "source.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), f"ИНН {_INN}", fontsize=12)
    doc.set_metadata({"author": _AUTHOR})
    doc.save(str(path))
    doc.close()

    parts = {part.name: part for part in pdf_parts(path)}
    assert "page 1" in parts
    assert _INN in parts["page 1"].text
    assert _INN.encode("utf-8") in parts["page 1"].raw
    assert "metadata" in parts
    assert _AUTHOR in parts["metadata"].text
    assert _AUTHOR.encode("utf-8") in parts["metadata"].raw


def test_pdf_parts_are_one_per_page_plus_metadata(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "multipage.pdf"
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 100), f"страница {i + 1}", fontsize=12)
    doc.save(str(path))
    doc.close()

    names = [part.name for part in pdf_parts(path)]
    assert names == ["page 1", "page 2", "page 3", "metadata"]
