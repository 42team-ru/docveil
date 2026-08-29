"""Генератор синтетического размеченного корпуса.

Реальные документы от заказчика придут позже и лягут в `fixtures/real/`
(он в .gitignore). До тех пор метрики считаются на синтетике, которая
намеренно содержит трудные случаи, а не только удобные.

Разметка пишется рядом в `<имя>.labels.json` — её формат читает `masker.eval`.
"""

from __future__ import annotations

import json
import pathlib

from docx import Document as DocxDocument
from docx.shared import Pt

OUT = pathlib.Path(__file__).resolve().parents[1] / "fixtures" / "labeled"


def contract_01() -> tuple[DocxDocument, list[dict[str, str]]]:
    """Обычный договор поставки: две стороны, блоки реквизитов, таблица."""
    doc = DocxDocument()
    doc.core_properties.author = "Петрова Мария Сергеевна"
    doc.core_properties.title = "Договор поставки 44/2026"

    doc.add_heading("ДОГОВОР ПОСТАВКИ № 44/2026", level=1)
    doc.add_paragraph("г. Воронеж, 15 января 2026 г.")
    doc.add_paragraph(
        "Акционерное общество «Триема», ИНН 3662103003, КПП 366201001, "
        "ОГРН 1023601546902, именуемое в дальнейшем «Поставщик», в лице "
        "генерального директора Иванова Ивана Ивановича, с одной стороны, и"
    )
    doc.add_paragraph(
        "Общество с ограниченной ответственностью «Вектор», ИНН 7707083893, "
        "КПП 770701001, именуемое в дальнейшем «Покупатель», в лице директора "
        "Сидоровой Анны Петровны, с другой стороны, заключили договор о нижеследующем."
    )

    doc.add_heading("1. Реквизиты Поставщика", level=2)
    doc.add_paragraph("Адрес: 394018, г. Воронеж, ул. Кирова, д. 4, оф. 12")
    doc.add_paragraph("Расчётный счёт: 40702810100000000001")
    doc.add_paragraph("БИК: 042007681")
    doc.add_paragraph("Телефон: +7 (473) 250-10-10, e-mail: info@triema.example")

    doc.add_heading("2. Реквизиты Покупателя", level=2)
    doc.add_paragraph("Адрес: 101000, г. Москва, ул. Мясницкая, д. 26")
    doc.add_paragraph("Телефон: 8-910-347-51-07")
    doc.add_paragraph("E-mail: zakupki@vektor.example")

    doc.add_heading("3. Спецификация", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for i, row in enumerate(
        [
            ("Наименование", "Кол-во", "Цена, руб."),
            ("Резистор МЛТ-0,25", "1000", "12 500,00"),
            ("Конденсатор К73-17", "500", "8 300,00"),
        ]
    ):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val

    doc.add_heading("4. Подписи", level=2)
    doc.add_paragraph("Поставщик: ______________ И.И. Иванов")
    doc.add_paragraph("Покупатель: ______________ А.П. Сидорова")

    labels = [
        {"type": "org_name", "text": "Акционерное общество «Триема»", "party": "supplier"},
        {"type": "inn", "text": "3662103003", "party": "supplier"},
        {"type": "kpp", "text": "366201001", "party": "supplier"},
        {"type": "ogrn", "text": "1023601546902", "party": "supplier"},
        {"type": "person", "text": "Иванова Ивана Ивановича", "party": "supplier"},
        {"type": "bank_account", "text": "40702810100000000001", "party": "supplier"},
        {"type": "bik", "text": "042007681", "party": "supplier"},
        {"type": "phone", "text": "+7 (473) 250-10-10", "party": "supplier"},
        {"type": "email", "text": "info@triema.example", "party": "supplier"},
        {"type": "org_name", "text": "Общество с ограниченной ответственностью «Вектор»", "party": "buyer"},
        {"type": "inn", "text": "7707083893", "party": "buyer"},
        {"type": "kpp", "text": "770701001", "party": "buyer"},
        {"type": "person", "text": "Сидоровой Анны Петровны", "party": "buyer"},
        {"type": "phone", "text": "8-910-347-51-07", "party": "buyer"},
        {"type": "email", "text": "zakupki@vektor.example", "party": "buyer"},
    ]
    return doc, labels


def contract_02_hard() -> tuple[DocxDocument, list[dict[str, str]]]:
    """Трудный случай: ИНН физлица, ловушки-непроходящие-контрольную-сумму.

    Здесь проверяется, что детектор не ловится на числа той же длины:
    номер накладной из 10 цифр и «ОГРН» из 13 цифр с битой контрольной
    цифрой обязаны быть пропущены. Регулярка без контрольной суммы
    провалит именно этот файл.
    """
    doc = DocxDocument()
    doc.core_properties.author = "Козлов П.А."

    doc.add_heading("АКТ СВЕРКИ ВЗАИМНЫХ РАСЧЁТОВ", level=1)
    doc.add_paragraph(
        "Индивидуальный предприниматель Кузнецов Пётр Алексеевич, "
        "ИНН 500100732259, СНИЛС 112-233-445 95."
    )
    doc.add_paragraph("Накладная № 3662103004 от 12.02.2026.")   # ловушка: битый ИНН
    doc.add_paragraph("Внутренний код операции: 1023601546903.")  # ловушка: битый ОГРН
    doc.add_paragraph("Паспорт 20 04 123456, выдан 10.05.2018.")
    doc.add_paragraph("Контакт: pkuznetsov@example.org")

    labels = [
        {"type": "person", "text": "Кузнецов Пётр Алексеевич", "party": "third_party"},
        {"type": "inn", "text": "500100732259", "party": "third_party"},
        {"type": "snils", "text": "112-233-445 95", "party": "third_party"},
        {"type": "passport", "text": "20 04 123456", "party": "third_party"},
        {"type": "email", "text": "pkuznetsov@example.org", "party": "third_party"},
    ]
    return doc, labels


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, builder in (("contract_01", contract_01), ("contract_02_hard", contract_02_hard)):
        doc, labels = builder()
        for p in doc.paragraphs:
            for run in p.runs:
                run.font.size = run.font.size or Pt(11)
        doc.save(OUT / f"{name}.docx")
        (OUT / f"{name}.labels.json").write_text(
            json.dumps({"entities": labels}, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"{name}.docx — сущностей в разметке: {len(labels)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
