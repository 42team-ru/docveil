"""Разбор DOCX в `Document`. T1.1.

Сегмент — непустой абзац основного текста (`w:body/w:p`). Якорь —
`("body", para_idx)`, где `para_idx` — физический индекс абзаца в
`doc.paragraphs`: пропущенные пустые абзацы физически существуют, и рендер
обязан находить абзац по индексу, а не по счётчику непустых.

Номер run'а в якорь не входит намеренно: сущность может пересекать границу
run'ов («ИНН 36» + «62103003» — обычное дело после правок в Word), и одним
индексом она не адресуется. Рендер получает по якорю абзац и сам вычисляет
покрывающие run'ы по `Entity.start/end` через `iter_runs`.

Таблицы (T1.11), колонтитулы и сноски (T1.12) в этот разбор не входят.
"""

from __future__ import annotations

import pathlib

from docx import Document as open_docx
from docx.oxml.text.run import CT_R
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from masker.model import Anchor, Document, Segment

#: Часть документа, которую обходим. Первый элемент якоря — дискриминатор:
#: T1.11/T1.12 добавят "table", "header", "footer", не трогая уже выданные якоря.
PART_BODY = "body"


def iter_runs(paragraph: Paragraph) -> list[Run]:
    """Runs абзаца в порядке документа, включая runs внутри гиперссылок.

    `paragraph.runs` отдаёт только прямых детей `w:p` и теряет текст
    гиперссылок, а `paragraph.text` их включает. Единственная форма, при
    которой смещения сегмента совпадают с обходом рендера, — эта:
    `"".join(r.text for r in iter_runs(p)) == p.text`.
    """
    elements: list[CT_R] = paragraph._p.xpath("./w:r | ./w:hyperlink/w:r")
    return [Run(r, paragraph) for r in elements]


def ingest_docx(path: str | pathlib.Path) -> Document:
    """Разобрать `.docx` в `Document` с сегментами и якорями."""
    path = pathlib.Path(path)
    doc = open_docx(str(path))

    segments: list[Segment] = []
    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text
        if not text.strip():
            continue
        segments.append(
            Segment(
                text=text,
                anchor=Anchor(
                    fmt="docx",
                    locator=(PART_BODY, para_idx),
                    label=f"абзац {para_idx + 1}",
                ),
                order=len(segments),
            )
        )

    props = doc.core_properties
    meta = {
        key: value
        for key, value in (
            ("author", props.author),
            ("last_modified_by", props.last_modified_by),
            ("title", props.title),
            ("subject", props.subject),
            ("comments", props.comments),
            ("category", props.category),
            ("keywords", props.keywords),
        )
        if value
    }

    return Document(path=str(path), fmt="docx", segments=segments, meta=meta)
