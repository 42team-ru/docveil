"""Тесты `masker.eval` — метрики по корпусу (T2.2.1, шаги 2 и 3).

Пайплайн (`masker.pipeline.mask_and_validate`) и профилирование
(`masker.eval._profile_judge_metrics`) в этих тестах подменены синтетикой:
цель — проверить логику самого `eval.py` (схлопывание пробелов, разрез по
форматам, гейт на утечки и дубли маркеров), а не гонять NER/LLM/рендер ещё
раз — это уже делает `make eval` отдельно и небыстро (полный корпус,
реальные PDF/DOCX).
"""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any

import pytest
from docx import Document as WordDocument

import masker.eval as eval_module
from masker.model import (
    Anchor,
    ArtifactLayout,
    Entity,
    EntityType,
    Leak,
    MaskGroup,
    MaskPlan,
    Replacement,
    Source,
    ValidationReport,
)
from masker.pipeline import MaskResult
from masker.run import RunFailedError

#: Заглушка `_profile_judge_metrics`: тесты этого файла не проверяют профили
#: и судью — только то, что `run()` строит из `load_corpus`/`mask_and_validate`.
_NEUTRAL_PROFILE_JUDGE_METRICS = {
    "cluster_purity": 1.0,
    "role_coverage": 1.0,
    "role_accuracy": 1.0,
    "critical_in_questions": 0.0,
    "questions_per_document": 0.0,
    "policy_questions_per_document": 0.0,
    "critical_unmasked": 0.0,
}

_EMPTY_VALIDATION = ValidationReport(
    leaked=(), residual=(), checked_artifacts=(), checked_parts=(), ok=True
)


def _replacement(entity_type: EntityType, text: str, *, ref: str = "R1") -> Replacement:
    entity = Entity(
        type=entity_type,
        text=text,
        segment_order=0,
        start=0,
        end=len(text),
        source=Source.RULE,
    )
    return Replacement(
        ref=ref,
        entity=entity,
        marker="[X]",
        group_id="G1",
        profile_id="",
        anchor=Anchor(fmt="docx", locator=("body", 0)),
    )


def _patch_common(
    monkeypatch: pytest.MonkeyPatch, corpus: list[tuple[Path, dict[str, Any]]]
) -> None:
    monkeypatch.setattr(eval_module, "load_corpus", lambda: corpus)
    monkeypatch.setattr(
        eval_module, "_profile_judge_metrics", lambda _corpus: dict(_NEUTRAL_PROFILE_JUDGE_METRICS)
    )


def _row(output: str, prefix: str) -> str:
    line = next((line for line in output.splitlines() if line.startswith(prefix)), None)
    assert line is not None, f"строка {prefix!r} не найдена в выводе:\n{output}"
    return line


def _patch_mask_and_validate(
    monkeypatch: pytest.MonkeyPatch, by_path: dict[str, MaskResult]
) -> None:
    """Подменить `masker.pipeline.mask_and_validate` на заранее заготовленные результаты."""

    @contextmanager
    def fake(path: Path, *, types: Any) -> Iterator[MaskResult]:
        yield by_path[str(path)]

    monkeypatch.setattr("masker.pipeline.mask_and_validate", fake)


def test_eval_collapses_whitespace_in_labels(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Разметка с лишними/непохожими пробелами не должна давать FN.

    Разметка PDF-корпуса пишется в «естественной однострочной форме» (схема
    разметки, пункт 1) и не обязана совпадать с тем, как ingest режет текст
    на сегменты. Без схлопывания пробелов двойной пробел в `text` разметки
    увёл бы это значение в false negative, хотя оно найдено верно.
    """
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    # Разметка — с двойным пробелом; найденная сущность — с одним пробелом,
    # ровно тот же класс расхождения, что описан в схеме разметки PDF.
    labels = {"entities": [{"type": "org_name", "text": "ООО  «Ромашка»"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.ORG_NAME, "ООО «Ромашка»"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    eval_module.run(gate=False)
    output = capsys.readouterr().out
    row = _row(output, "org_name")
    _, _precision, recall, _f1, fn, _fp = row.split()
    assert recall == "1.000", f"ожидали recall 1.0 при схлопывании пробелов, строка: {row!r}"
    assert fn == "0", f"двойной пробел в разметке не должен давать FN, строка: {row!r}"


def test_eval_reports_pdf_format_row(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Секция «ФОРМАТЫ» обязана появиться и содержать строку `pdf`.

    До шага 2 `_profile_judge_metrics` пропускал всё, что не `.docx`
    (`if path.suffix != ".docx": continue` — Д7 плана T2.2.1), а `run()` не
    строил разреза по форматам вовсе — идеальные цифры по DOCX маскировали
    провал по PDF.
    """
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"")
    corpus = [
        (docx_path, {"entities": [{"type": "inn", "text": "1234567890"}]}),
        (pdf_path, {"entities": [{"type": "inn", "text": "1234567890"}]}),
    ]
    _patch_common(monkeypatch, corpus)
    # PDF нарочно не находит ничего — именно так выглядит Д7 сегодня: PDF
    # физически участвует в прогоне, но детекция по нему хромает.
    docx_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    pdf_result = MaskResult(
        plan=MaskPlan(replacements=(), groups=(), skipped=(), requested_types=()),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): docx_result, str(pdf_path): pdf_result})

    eval_module.run(gate=False)
    output = capsys.readouterr().out
    assert "ФОРМАТЫ" in output
    docx_row = _row(output, "docx")
    pdf_row = _row(output, "pdf")
    assert docx_row.split()[2] == "1.000", f"docx recall должен остаться 1.0: {docx_row!r}"
    assert pdf_row.split()[2] == "0.000", f"pdf recall должен показать провал: {pdf_row!r}"


def test_eval_gate_fails_on_leak(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Синтетический план с заведомой утечкой — ворота обязаны провалиться.

    До шага 3 `leaked_total` нигде не читался: `ValidateAgent` уже находил
    утечки (T1.8), но `eval.py` их не суммировал и не проверял порогом.
    """
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    leak = Leak(
        kind="raw",
        artifact="masked_black.docx",
        part="word/document.xml",
        entity_type="inn",
        value="1234567890",
        ref="R1",
        group_id="G1",
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(leak,),
            residual=(),
            checked_artifacts=("masked_black.docx",),
            checked_parts=(),
            ok=False,
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "утечка обязана провалить ворота"
    assert _row(output, "leaked_total").split()[-1] == "1"


def test_eval_gate_passes_without_leak(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Симметричный случай: без утечки ворота по `leaked_total` не падают."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert _row(output, "leaked_total").split()[-1] == "0"


def test_eval_gate_fails_on_layout_loss(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Прямоугольник редакции стёр текст вне своих замен (Д10, план T2.2.2,
    шаг 5) — ворота обязаны провалиться на `layout_removed_chars`."""
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(pdf_path, labels)])
    layout = ArtifactLayout(
        artifact="masked_black.pdf",
        removed_chars=46,
        inserted_chars=0,
        pages=(0,),
        first_diff="стр. 1: ожидалось …Устава, с одной…, получено …Устава…",
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(),
            residual=(),
            checked_artifacts=("masked_black.pdf",),
            checked_parts=(),
            ok=True,
            layout=(layout,),
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(pdf_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "потеря текста вне замены обязана провалить ворота"
    assert _row(output, "layout_removed_chars").split()[-1] == "46"
    assert "masked_black.pdf" in output


def test_eval_gate_passes_without_layout_loss(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Симметричный случай: без потери вёрстки ворота по
    `layout_removed_chars` не падают."""
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(pdf_path, labels)])
    layout = ArtifactLayout(
        artifact="masked_black.pdf",
        removed_chars=0,
        inserted_chars=0,
        pages=(),
        first_diff="",
    )
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=ValidationReport(
            leaked=(),
            residual=(),
            checked_artifacts=("masked_black.pdf",),
            checked_parts=(),
            ok=True,
            layout=(layout,),
        ),
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(pdf_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert _row(output, "layout_removed_chars").split()[-1] == "0"


def test_eval_gate_fails_on_duplicate_marker(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Маркер, вставленный дважды на один `Replacement` (Д1), обязан провалить ворота."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "org_name", "text": "ООО «Ромашка»"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])

    artifact = tmp_path / "masked_black.docx"
    word_doc = WordDocument()
    # Один Replacement, но маркер физически вставлен в артефакт дважды —
    # ровно дефект Д1 (рендер искал строку заново и попал в неё повторно).
    word_doc.add_paragraph("[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]")
    word_doc.add_paragraph("[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]")
    word_doc.save(artifact)

    plan = MaskPlan(
        replacements=(_replacement(EntityType.ORG_NAME, "ООО «Ромашка»"),),
        groups=(
            MaskGroup(
                id="G1",
                key="org:1",
                type=EntityType.ORG_NAME,
                marker="[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]",
                profile_id="P1",
                role_label="ПОСТАВЩИК",
                number=1,
                refs=("R1",),
                sample="ООО «Ромашка»",
            ),
        ),
        skipped=(),
        requested_types=(),
    )
    result = MaskResult(plan=plan, validation=_EMPTY_VALIDATION, artifacts=(artifact,))
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 1, "дубль маркера обязан провалить ворота"
    assert _row(output, "duplicate_markers").split()[-1] == "1"


def test_duplicate_marker_count_ignores_exact_match(tmp_path: Path) -> None:
    """Маркер, вставленный ровно по разу на `Replacement`, — не дубль."""
    artifact = tmp_path / "masked_black.docx"
    word_doc = WordDocument()
    word_doc.add_paragraph("[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]")
    word_doc.save(artifact)
    plan = MaskPlan(
        replacements=(),
        groups=(
            MaskGroup(
                id="G1",
                key="org:1",
                type=EntityType.ORG_NAME,
                marker="[ПОСТАВЩИК-ОРГАНИЗАЦИЯ]",
                profile_id="P1",
                role_label="ПОСТАВЩИК",
                number=1,
                refs=("R1",),
                sample="ООО «Ромашка»",
            ),
        ),
        skipped=(),
        requested_types=(),
    )
    assert eval_module.duplicate_marker_count(plan, (artifact,)) == 0


def test_eval_gate_fails_on_render_failure_but_keeps_measuring_the_rest(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Прогон корпуса — измерительный инструмент: аномальная сущность на
    одном документе не имеет права ослепить ворота целиком и скрыть
    метрики по остальным документам (план T2.2.1, пачка 4). При этом
    падение обязано остаться видимым — новой метрикой с адресом проблемы,
    а не тихим пропуском документа.
    """
    broken_path = tmp_path / "broken.pdf"
    broken_path.write_bytes(b"")
    healthy_path = tmp_path / "healthy.docx"
    healthy_path.write_bytes(b"")
    corpus = [
        (broken_path, {"entities": [{"type": "org_name", "text": "ГО и ЧС"}]}),
        (healthy_path, {"entities": [{"type": "inn", "text": "1234567890"}]}),
    ]
    _patch_common(monkeypatch, corpus)
    healthy_result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )

    @contextmanager
    def fake(path: Path, *, types: Any) -> Iterator[MaskResult]:
        if str(path) == str(broken_path):
            cause = ValueError(
                "маркер '[СТОРОНА-27-ОРГАНИЗАЦИЯ]' не помещается в прямоугольник "
                "ни при одном размере шрифта"
            )
            raise RunFailedError("thread123", "render", cause)
        yield healthy_result

    monkeypatch.setattr("masker.pipeline.mask_and_validate", fake)

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out

    assert code == 1, "падение рендера обязано провалить ворота"
    assert _row(output, "render_failures").split()[-1] == "1"
    assert broken_path.name in output, "имя упавшего документа обязано быть в отчёте"
    assert "СТОРОНА-27-ОРГАНИЗАЦИЯ" in output, "маркер, на котором упал рендер, обязан быть виден"
    # Здоровый документ не должен пострадать от падения соседнего.
    inn_row = _row(output, "inn")
    assert inn_row.split()[2] == "1.000", (
        f"метрика по healthy.docx обязана посчитаться: {inn_row!r}"
    )


def test_eval_gate_passes_without_render_failure(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    """Симметричный случай: без падений рендера ворота по `render_failures` не падают."""
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"")
    labels = {"entities": [{"type": "inn", "text": "1234567890"}]}
    _patch_common(monkeypatch, [(docx_path, labels)])
    result = MaskResult(
        plan=MaskPlan(
            replacements=(_replacement(EntityType.INN, "1234567890"),),
            groups=(),
            skipped=(),
            requested_types=(),
        ),
        validation=_EMPTY_VALIDATION,
        artifacts=(),
    )
    _patch_mask_and_validate(monkeypatch, {str(docx_path): result})

    code = eval_module.run(gate=True)
    output = capsys.readouterr().out
    assert code == 0
    assert _row(output, "render_failures").split()[-1] == "0"
