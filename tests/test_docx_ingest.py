"""Контрактные тесты ingest DOCX: только основной текст и core properties."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document as open_docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import _BaseHeaderFooter
from docx.text.paragraph import Paragraph

from masker.detect.rules import detect_by_rules
from masker.ingest.docx_ingest import ingest_docx, iter_runs
from masker.model import EntityType

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures" / "labeled"
CONTRACTS = [
    (FIXTURES / "contract_01.docx", 17),
    (FIXTURES / "contract_02_hard.docx", 6),
]


def append_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    """Добавить внешнюю гиперссылку, как она будет сохранена в DOCX."""
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


@pytest.mark.parametrize(("path", "expected_count"), CONTRACTS)
def test_segment_count_equals_nonempty_body_paragraphs(path: Path, expected_count: int) -> None:
    """Сегменты совпадают с непустыми абзацами тела, но не с ячейками таблиц."""
    expected = sum(bool(paragraph.text.strip()) for paragraph in open_docx(path).paragraphs)

    document = ingest_docx(path)

    assert expected == expected_count
    assert len(document.segments) == expected == expected_count


def test_order_is_dense_from_zero() -> None:
    document = ingest_docx(CONTRACTS[0][0])
    print(document)

    assert [segment.order for segment in document.segments] == list(range(len(document.segments)))


def test_anchor_keeps_physical_paragraph_index() -> None:
    source = open_docx(CONTRACTS[0][0])
    document = ingest_docx(CONTRACTS[0][0])

    for segment in document.segments:
        part, para_idx = segment.anchor.locator
        assert part == "body"
        assert segment.anchor.fmt == "docx"
        assert segment.anchor.label == f"абзац {para_idx + 1}"
        assert source.paragraphs[para_idx].text == segment.text


def test_segment_text_is_charwise_equal_to_paragraph_text() -> None:
    source = open_docx(CONTRACTS[0][0])
    document = ingest_docx(CONTRACTS[0][0])

    for segment in document.segments:
        para_idx = segment.anchor.locator[1]
        assert isinstance(para_idx, int)
        paragraph = source.paragraphs[para_idx]
        assert segment.text == paragraph.text
        assert "".join(run.text for run in iter_runs(paragraph)) == segment.text


def test_table_text_is_not_ingested_yet() -> None:
    document = ingest_docx(CONTRACTS[0][0])

    assert all("Резистор МЛТ-0,25" not in segment.text for segment in document.segments)
    assert "Резистор МЛТ-0,25" not in document.text()


def test_meta_from_core_properties() -> None:
    document = ingest_docx(CONTRACTS[0][0])

    assert document.meta["author"] == "Петрова Мария Сергеевна"
    assert document.meta["title"]
    assert "last_modified_by" not in document.meta
    assert "subject" not in document.meta
    assert "category" not in document.meta
    assert all(isinstance(value, str) for value in document.meta.values())


def test_document_text_joins_segments_with_newline() -> None:
    document = ingest_docx(CONTRACTS[0][0])

    assert document.text() == "\n".join(segment.text for segment in document.segments)
    assert "ИНН 3662103003" in document.text()


def test_ingest_does_not_modify_source_file(tmp_path: Path) -> None:
    path = tmp_path / "contract.docx"
    path.write_bytes(CONTRACTS[0][0].read_bytes())
    before = hashlib.sha256(path.read_bytes()).digest()

    ingest_docx(path)

    assert hashlib.sha256(path.read_bytes()).digest() == before


def test_ingest_never_reads_headers_or_footers() -> None:
    with patch.object(
        _BaseHeaderFooter,
        "_get_or_add_definition",
        side_effect=AssertionError("ingest не должен читать колонтитул"),
    ):
        ingest_docx(CONTRACTS[0][0])


def test_ingest_is_deterministic() -> None:
    path = CONTRACTS[0][0]
    expected = [(s.order, s.anchor.locator, s.text) for s in ingest_docx(path).segments]
    repeated = [(s.order, s.anchor.locator, s.text) for s in ingest_docx(path).segments]
    script = (
        "import json, sys; "
        f"sys.path.insert(0, {str(ROOT / 'src')!r}); "
        "from masker.ingest.docx_ingest import ingest_docx; "
        f"d = ingest_docx({str(path)!r}); "
        "print(json.dumps([(s.order, s.anchor.locator, s.text) for s in d.segments], "
        "ensure_ascii=False))"
    )
    result = subprocess.run(
        [sys.executable, "-c", script],
        check=True,
        capture_output=True,
        env={**os.environ, "PYTHONHASHSEED": "1"},
        text=True,
    )

    assert repeated == expected
    assert json.loads(result.stdout) == [
        [order, list(locator), text] for order, locator, text in expected
    ]


def test_runs_are_glued_without_separator(tmp_path: Path) -> None:
    path = tmp_path / "split-runs.docx"
    source = open_docx()
    paragraph = source.add_paragraph()
    for text in ("ИНН ", "36621", "03003"):
        paragraph.add_run(text)
    source.save(path)

    segment = ingest_docx(path).segments[0]
    entities = detect_by_rules([segment])

    assert segment.text == "ИНН 3662103003"
    assert any(entity.type == EntityType.INN for entity in entities)


def test_hyperlink_text_is_ingested_and_reachable_by_runs(tmp_path: Path) -> None:
    path = tmp_path / "hyperlink.docx"
    source = open_docx()
    paragraph = source.add_paragraph("Почта: ")
    append_hyperlink(paragraph, "info@triema.example", "mailto:info@triema.example")
    paragraph.add_run(" — пишите.")
    source.save(path)

    document = ingest_docx(path)
    reloaded_paragraph = open_docx(path).paragraphs[0]
    runs = iter_runs(reloaded_paragraph)

    assert "info@triema.example" in document.segments[0].text
    assert sum(len(run.text) for run in runs) == len(document.segments[0].text)
    assert len(runs) > len(reloaded_paragraph.runs)


def test_empty_paragraphs_are_skipped_but_indexes_stay_physical(tmp_path: Path) -> None:
    path = tmp_path / "empty-paragraphs.docx"
    source = open_docx()
    source.add_paragraph("")
    source.add_paragraph("Текст A")
    source.add_paragraph("   ")
    source.add_paragraph("Текст B")
    source.save(path)

    document = ingest_docx(path)

    assert [segment.order for segment in document.segments] == [0, 1]
    assert [segment.anchor.locator for segment in document.segments] == [("body", 1), ("body", 3)]


def test_document_without_core_properties_gives_empty_meta(tmp_path: Path) -> None:
    path = tmp_path / "empty-document.docx"
    source = open_docx()
    properties = source.core_properties
    for name in (
        "author",
        "last_modified_by",
        "title",
        "subject",
        "comments",
        "category",
        "keywords",
    ):
        setattr(properties, name, "")
    source.save(path)

    document = ingest_docx(path)

    assert document.meta == {}
    assert document.segments == []
    assert document.text() == ""
